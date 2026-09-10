"""Audited XLSX exports from one materialized, tenant-scoped evidence selection."""

import hashlib
import io
import json
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from uuid import uuid4

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import String, cast, literal, select

from app.core.database import set_tenant_context
from app.models.transaction_ops import TransactionCase
from app.services.audit_service import log_event
from app.services.transaction_ops.case_groups import _pattern, _signature
from app.services.transaction_ops.state_service import StateError
from app.services.transaction_ops.workspace_results import filtered_query, selected_evidence

MAX_EXPORT_ROWS = 50000
REVIEW_STATUSES = {"difference", "mismatch", "missing_in_netsuite", "ambiguous", "currency_mismatch"}


def _text(cell, value):
    # Untrusted references/labels must never become spreadsheet formulas.
    cell.value = str(value) if value is not None else "—"
    cell.data_type = "s"


def _money(cell, value):
    try:
        if not isinstance(value, str) or len(value) > 100:
            raise ValueError
        amount = Decimal(value)
        if not amount.is_finite() or abs(amount.adjusted()) > 100:
            raise ValueError
    except (InvalidOperation, ValueError):
        _text(cell, "—")
        return
    # Excel supports only 15 significant digits. Preserve larger values as
    # exact text instead of silently changing a financial amount.
    if len(amount.as_tuple().digits) > 15 or -amount.as_tuple().exponent > 12:
        _text(cell, value)
    else:
        cell.value = amount
        places = min(12, max(2, -amount.as_tuple().exponent))
        cell.number_format = "#,##0." + "0" * places + ";[Red]-#,##0." + "0" * places + ";0." + "0" * places
    cell.alignment = Alignment(horizontal="right")


def _sheet(wb, name, columns, rows, *, money_columns=(), numeric_columns=()):
    ws = wb.create_sheet(name)
    ws.append(columns)
    for cell in ws[1]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="C94213")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 34
    for index, row in enumerate(rows, start=2):
        for col, value in enumerate(row, start=1):
            cell = ws.cell(index, col)
            if col in money_columns:
                _money(cell, value)
            elif col in numeric_columns and type(value) is int:
                cell.value = value
            else:
                _text(cell, value)
            cell.font = Font(name="Arial", size=10)
            if index % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F4F4F5")
        ws.row_dimensions[index].height = 20
    for index, header in enumerate(columns, start=1):
        width = max(len(header), *(len(str(row[index - 1] or "")) for row in rows[:100])) if rows else len(header)
        ws.column_dimensions[get_column_letter(index)].width = min(48, max(16, width + 2))
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    ws.print_title_rows = "1:1"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    return ws


def build_workbook(rows, scopes, *, status, search, generated_at, report_id):
    scope_map = {item["run_id"]: item for item in scopes}
    details, groups = [], {}
    for row in rows:
        scope = scope_map[row["review_run_id"]]
        config = scope["config"]
        balance = row["balance"] or {}
        category = (
            "matched"
            if row["status"] == "matched"
            else "needs_review"
            if row["status"] in REVIEW_STATUSES
            else "not_verified"
        )
        group_category = status or "needs_review"
        grouped = category == group_category and group_category != "matched" and row["scope"] is not None
        group_id = row["group_id"] if grouped else None
        name = config.get("name") or f"Entity {config.get('subsidiary_id', 'unavailable')}"
        if grouped:
            if group_id not in groups:
                groups[group_id] = [group_id, _pattern(row), name, row["currency"], row["target_state"], 0]
            groups[group_id][-1] += 1
        amounts = balance.get("amounts") or {}
        details.append(
            [
                row["order_reference"],
                name,
                row["currency"],
                row["target_currency"],
                category,
                *[
                    (amounts.get(metric) or {}).get(side)
                    for metric in ("order_total", "tax", "refunds")
                    for side in ("source", "target", "delta")
                ],
                row["status"],
                row["target_state"],
                group_id,
                row["case_id"],
                str(row["run_id"]),
                row["updated_at"].isoformat(),
                row["review_run_id"],
            ]
        )
    wb = Workbook()
    wb.remove(wb.active)
    info = [
        ["Report", "Transaction reconciliation"],
        ["Report ID", report_id],
        ["Generated at (UTC)", generated_at.isoformat()],
        ["Status filter", status or "All results"],
        ["Order search", search or "All order references"],
        ["Exported orders", len(details)],
        ["Issue groups", len(groups)],
        ["Grouped orders", sum(group[-1] for group in groups.values())],
        ["Coverage", "Recorded evidence only. Exporting does not complete a period scan or certify settlement."],
        ["Variance", "Source minus ERP. Order total includes tax; do not add tax variance to order variance."],
        ["Amounts", "Every nonzero delta is retained. Unknown = —. Values exceeding Excel precision stay exact text."],
        ["Currency", "Amounts are in each row's currency; different currencies are never totaled together."],
        [
            "Group scope",
            "All results groups Needs review; otherwise the selected status. Matched orders have no issue groups.",
        ],
        [
            "Approval",
            "This report does not approve financial changes. Group membership is not an execution authorization.",
        ],
    ]
    for scope in scopes:
        config = scope["config"]
        info.append(
            [
                config.get("name") or f"Entity {config.get('subsidiary_id')}",
                f"{scope['period']['start']} to {scope['period']['end']} (exclusive); "
                f"source date basis: {scope.get('window_basis', 'completed_at')}; review run {scope['run_id']}",
            ]
        )
    overview = _sheet(wb, "Report", ["Field", "Value"], info, numeric_columns=(2,))
    overview.column_dimensions["A"].width = 28
    overview.column_dimensions["B"].width = 110
    overview.freeze_panes = "B2"
    overview.auto_filter.ref = None
    for row in overview.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="center")
        overview.row_dimensions[row[0].row].height = 32
    columns = ["Order number", "Entity", "Source currency", "ERP currency", "Finding"]
    columns += [
        f"{metric} — {side}"
        for metric in ("Order total", "VAT / tax", "Refunds")
        for side in ("Source", "ERP", "Variance")
    ]
    columns += [
        "Evidence status",
        "ERP lifecycle",
        "Issue group ID",
        "Case ID",
        "Evidence run ID",
        "Observed at (UTC)",
        "Review run ID",
    ]
    _sheet(wb, "Reconciliation", columns, details, money_columns=range(6, 15))
    _sheet(
        wb,
        "Issue groups",
        ["Group ID", "Pattern", "Entity", "Currency", "ERP lifecycle", "Orders"],
        sorted(groups.values(), key=lambda row: (-row[-1], row[0])),
        numeric_columns=(6,),
    )
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue(), len(details), len(groups)


async def export_review(db, actor, run_ids, *, status=None, search=""):
    tenant_id, actor_id = actor.tenant_id, actor.id
    latest, scopes = await selected_evidence(db, tenant_id, run_ids)
    filtered = filtered_query(latest, status, search).subquery()
    source = (
        select(
            filtered,
            filtered.c.report_json.label("latest_report_json"),
            TransactionCase.scope_json,
            literal(str(tenant_id), String).label("tenant_id"),
        )
        .outerjoin(
            TransactionCase,
            (cast(TransactionCase.id, String) == filtered.c.report_json["case_id"].astext)
            & (TransactionCase.tenant_id == tenant_id)
            & (TransactionCase.order_reference == filtered.c.order_reference),
        )
        .subquery()
    )
    scope_key = json.dumps(
        [sorted({str(item["run_id"]) for item in scopes}), status or "needs_review", search], separators=(",", ":")
    )
    columns, identifier = _signature(source, scope_key)
    # _signature consumes the same report/scope as the issue-group endpoint.
    report_id = str(uuid4())
    payload = {
        "review_run_ids": [item["run_id"] for item in scopes],
        "status": status,
        "search": search,
        "financial_approval": None,
    }
    await log_event(
        db,
        tenant_id,
        category="transaction_ops",
        action="transaction_ops.report.requested",
        actor_id=actor_id,
        resource_type="transaction_report",
        resource_id=report_id,
        payload=payload,
    )
    await db.commit()
    await set_tenant_context(db, str(tenant_id))
    try:
        rows = (
            (
                await db.execute(
                    select(
                        source.c.order_reference,
                        source.c.review_run_id,
                        source.c.run_id,
                        source.c.updated_at,
                        source.c.report_json["balance"].label("balance"),
                        source.c.report_json["case_id"].astext.label("case_id"),
                        *columns,
                        identifier,
                    )
                    .order_by(source.c.order_reference, source.c.id)
                    .limit(MAX_EXPORT_ROWS + 1)
                )
            )
            .mappings()
            .all()
        )
        if len(rows) > MAX_EXPORT_ROWS:
            raise StateError("export_too_large_narrow_period_or_filters", 422)
        content, count, group_count = build_workbook(
            rows, scopes, status=status, search=search, generated_at=datetime.now(timezone.utc), report_id=report_id
        )
        await log_event(
            db,
            tenant_id,
            category="transaction_ops",
            action="transaction_ops.report.generated",
            actor_id=actor_id,
            resource_type="transaction_report",
            resource_id=report_id,
            payload={**payload, "rows": count, "groups": group_count, "sha256": hashlib.sha256(content).hexdigest()},
        )
        await db.commit()
        return content, report_id
    except Exception as exc:
        await db.rollback()
        await set_tenant_context(db, str(tenant_id))
        await log_event(
            db,
            tenant_id,
            category="transaction_ops",
            action="transaction_ops.report.failed",
            actor_id=actor_id,
            resource_type="transaction_report",
            resource_id=report_id,
            status="error",
            payload={**payload, "code": exc.code if isinstance(exc, StateError) else "export_failed"},
        )
        await db.commit()
        raise
