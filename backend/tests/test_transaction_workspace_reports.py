from copy import deepcopy
from datetime import timedelta
from io import BytesIO
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.api.v1.transaction_ops import router
from app.models.audit import AuditEvent
from app.models.transaction_ops import TransactionCase
from app.services.transaction_ops.case_groups import list_groups
from app.services.transaction_ops.workspace_results import record_page, review_page
from tests.conftest import enable_feature_flag
from tests.test_transaction_review_results import evidence, recheck
from tests.test_transaction_review_slices import review


@pytest.fixture(autouse=True)
async def routes(app):
    app.include_router(router, prefix="/api/v1")


async def linked(db, actor, root, case, *, status="difference", at=None):
    when = at or root.created_at
    row = await evidence(db, actor, root, case.order_reference, status, when, source_id=str(case.id), currency="EUR")
    report = deepcopy(case.latest_report_json)
    report.update(case_id=str(case.id), source={"record_id": str(case.id)})
    report["balance"]["status"] = status
    row.report_json = report
    row.updated_at = when + timedelta(microseconds=1)
    await db.flush()
    return row


async def fixture_rows(db, actor, monkeypatch, *, malicious=False, count=27):
    _, first = await review(db, actor, monkeypatch)
    _, second = await review(db, actor, monkeypatch)
    cases = []
    for index in range(count):
        identity = uuid4()
        root = first if index < count // 2 else second
        report = {
            "balance": {
                "status": "difference",
                "currency": "EUR",
                "target_currency": "EUR",
                "missing_metrics": [],
                "amounts": {
                    metric: {
                        "source": "0.00",
                        "target": "0.02" if metric == "tax" else "0.00",
                        "delta": "-0.02" if metric == "tax" else "0.00",
                    }
                    for metric in ("order_total", "tax", "refunds")
                },
            },
            "targets": [{"status": "fulfilled"}],
        }
        case = TransactionCase(
            id=identity,
            tenant_id=actor.tenant_id,
            case_key=identity.hex,
            order_reference='=HYPERLINK("https://example.invalid","bad")'
            if malicious and index == 0
            else f"R{index:09d}",
            scope_json={
                key: root.config_snapshot.get(key)
                for key in (
                    "source_connection_id",
                    "source_step_id",
                    "netsuite_account_id",
                    "subsidiary_id",
                    "record_type",
                )
            },
            status="open",
            first_observed_at=root.created_at,
            last_observed_at=root.created_at,
            latest_report_json=report,
        )
        db.add(case)
        await db.flush()
        await linked(db, actor, root, case)
        cases.append(case)
    return first, second, cases


async def test_combined_pagination_is_global_stable_and_rechecks_precede_filters(db, admin_user, monkeypatch):
    actor = admin_user[0]
    first, second, cases = await fixture_rows(db, actor, monkeypatch)
    ids = [first.id, second.id]
    page1 = await review_page(db, actor.tenant_id, ids, limit=25)
    page2 = await review_page(db, actor.tenant_id, ids, limit=25, offset=25)
    assert page1["total"] == 27 and len(page1["items"]) == 25 and page1["has_next"]
    assert len(page2["items"]) == 2 and not page2["has_next"]
    assert len({item["id"] for item in page1["items"] + page2["items"]}) == 27
    assert [item["order_reference"] for item in page1["items"] + page2["items"]] == sorted(
        c.order_reference for c in cases
    )
    assert (await review_page(db, actor.tenant_id, [first.id]))["total"] == 13
    assert (await review_page(db, actor.tenant_id, [first.id, first.id]))["total"] == 13
    later = await recheck(db, actor, first)
    await linked(db, actor, later, cases[0], status="matched", at=first.created_at + timedelta(seconds=2))
    filtered = await review_page(db, actor.tenant_id, ids, status="needs_review")
    assert filtered["total"] == 26 and filtered["summary"]["matched"] == 1
    assert (await review_page(db, actor.tenant_id, ids, search=cases[-1].order_reference))["total"] == 1
    assert (await review_page(db, actor.tenant_id, ids, offset=10000))["items"] == []


async def test_export_contains_all_pages_exact_amounts_safe_text_and_audit(client, db, admin_user, monkeypatch):
    actor, headers = admin_user
    first, second, cases = await fixture_rows(db, actor, monkeypatch, malicious=True, count=75)
    # Stored source labels/references are untrusted spreadsheet input.
    cases[0].latest_report_json = {
        **cases[0].latest_report_json,
        "balance": {
            **cases[0].latest_report_json["balance"],
            "amounts": {
                "order_total": {
                    "source": "123456789012345.123456",
                    "target": "0.00",
                    "delta": "123456789012345.123456",
                },
                "tax": {"source": "0.00", "target": "0.02", "delta": "-0.02"},
                "refunds": {"source": None, "target": "0.00", "delta": "1e-20"},
            },
        },
    }
    # Replace the original reference in the root fixture, keeping 75 cohort rows.
    from app.models.transaction_ops import TransactionFinding

    old = await db.scalar(
        select(TransactionFinding).where(
            TransactionFinding.run_id == first.id, TransactionFinding.order_reference == cases[0].order_reference
        )
    )
    old.report_json = {
        **deepcopy(cases[0].latest_report_json),
        "case_id": str(cases[0].id),
        "source": {"record_id": str(cases[0].id)},
    }
    await db.flush()
    response = await client.post(
        "/api/v1/transaction-ops/review-export",
        headers=headers,
        json={"review_run_ids": [str(first.id), str(second.id)], "status": "needs_review"},
    )
    assert response.status_code == 200, response.text[:200] if response.status_code != 200 else ""
    assert response.headers["cache-control"] == "no-store"
    wb = load_workbook(BytesIO(response.content), data_only=False)
    assert wb.sheetnames == ["Report", "Reconciliation", "Issue groups"]
    ws = wb["Reconciliation"]
    assert ws.max_row == 76 and ws.freeze_panes == "C2"
    assert ws["A2"].value.startswith("=HYPERLINK") and ws["A2"].data_type == "s"
    assert ws["F2"].value == "123456789012345.123456" and ws["F2"].data_type == "s"
    assert ws["K2"].value == -0.02 and ws["L2"].value == "—" and ws["M2"].value == 0
    assert ws["N2"].value == "1e-20"
    assert all(cell.data_type != "f" for sheet in wb for row in sheet for cell in row)
    assert sum(row[5].value for row in wb["Issue groups"].iter_rows(min_row=2)) == 75
    events = (
        await db.scalars(
            select(AuditEvent).where(
                AuditEvent.tenant_id == actor.tenant_id, AuditEvent.action == "transaction_ops.report.generated"
            )
        )
    ).all()
    assert len(events) == 1 and events[0].actor_id == actor.id and events[0].payload["rows"] == 75
    import hashlib

    assert events[0].payload["sha256"] == hashlib.sha256(response.content).hexdigest()
    assert events[0].payload["financial_approval"] is None


async def test_pages_export_and_group_totals_keep_tenant_auth_scope_and_bounds(
    client, db, admin_user, admin_user_b, monkeypatch
):
    actor, headers = admin_user
    first, second, cases = await fixture_rows(db, actor, monkeypatch)
    for flag in ("celigo", "reconciliation"):
        await enable_feature_flag(db, admin_user_b[0].tenant_id, flag)
    path = "/api/v1/transaction-ops/review-results"
    query = [("review_run_ids", str(first.id)), ("review_run_ids", str(second.id)), ("limit", "25")]
    assert (await client.get(path, params=query)).status_code == 401
    page = await client.get(path, params=query, headers=headers)
    assert page.status_code == 200 and len(page.json()["items"]) == 25
    assert (await client.get(path, params=query, headers=admin_user_b[1])).status_code == 404
    assert (
        await client.get(path, params={"review_run_ids": str(first.id), "limit": 501}, headers=headers)
    ).status_code == 422
    assert (
        await client.post(
            "/api/v1/transaction-ops/review-export", json={"review_run_ids": [str(first.id)]}, headers=admin_user_b[1]
        )
    ).status_code == 404
    assert (
        await client.post(
            "/api/v1/transaction-ops/review-export",
            json={"review_run_ids": [str(first.id)], "actor_id": str(uuid4())},
            headers=headers,
        )
    ).status_code == 422
    for view in ("cases", "runs", "proposals"):
        own = await record_page(db, actor.tenant_id, view, limit=1)
        foreign = await record_page(db, admin_user_b[0].tenant_id, view, limit=1)
        assert foreign["total"] == 0
        if view != "proposals":
            assert len(own["items"]) == 1 and own["has_next"]
    groups = await list_groups(db, actor.tenant_id, review_run_ids=[first.id, second.id], limit=1)
    assert groups["total_cases"] == 27
    beyond = await list_groups(db, actor.tenant_id, review_run_ids=[first.id, second.id], offset=1000)
    assert beyond["total_cases"] == 27 and beyond["total_groups"] == groups["total_groups"] and not beyond["groups"]


async def test_export_refuses_oversize_without_truncation_and_audits_failure(client, db, admin_user, monkeypatch):
    actor, headers = admin_user
    tenant_id, actor_id = actor.tenant_id, actor.id
    first, _, _ = await fixture_rows(db, actor, monkeypatch)
    monkeypatch.setattr("app.services.transaction_ops.excel_report.MAX_EXPORT_ROWS", 1)
    result = await client.post(
        "/api/v1/transaction-ops/review-export", json={"review_run_ids": [str(first.id)]}, headers=headers
    )
    assert result.status_code == 422 and "export_too_large" in result.text
    event = await db.scalar(
        select(AuditEvent).where(
            AuditEvent.tenant_id == tenant_id, AuditEvent.action == "transaction_ops.report.failed"
        )
    )
    assert event.actor_id == actor_id and event.status == "error"


async def test_page_sizes_50_100_500_have_no_hidden_entity_multiplier(client, db, admin_user, monkeypatch):
    actor, headers = admin_user
    first, second, _ = await fixture_rows(db, actor, monkeypatch, count=501)
    params = [("review_run_ids", str(first.id)), ("review_run_ids", str(second.id))]
    path = "/api/v1/transaction-ops/review-results"
    for size in (50, 100, 500):
        response = await client.get(path, params=[*params, ("limit", str(size))], headers=headers)
        assert response.status_code == 200
        assert response.json()["total"] == 501 and len(response.json()["items"]) == size
    tail = await client.get(path, params=[*params, ("limit", "500"), ("offset", "500")], headers=headers)
    assert len(tail.json()["items"]) == 1 and not tail.json()["has_next"]
    default = await client.get(path, params=params, headers=headers)
    assert len(default.json()["items"]) == 50
    for view in ("cases", "runs", "proposals"):
        page = await client.get(
            "/api/v1/transaction-ops/workspace-page", params={"view": view, "limit": 500}, headers=headers
        )
        assert page.status_code == 200
    groups = await client.get("/api/v1/transaction-ops/case-groups", params={"limit": 500}, headers=headers)
    assert groups.status_code == 200 and groups.json()["total_cases"] == 501


async def test_group_listing_can_page_more_than_fifty_groups(db, tenant_a):
    from tests.test_transaction_case_groups import seed

    for index in range(101):
        await seed(db, tenant_a.id, scope={"source_connection_id": "source", "subsidiary_id": str(index)})
    page = await list_groups(db, tenant_a.id, limit=100)
    assert len(page["groups"]) == 100 and page["total_groups"] == 101 and page["has_next"]
    tail = await list_groups(db, tenant_a.id, limit=100, offset=100)
    assert len(tail["groups"]) == 1 and not tail["has_next"]
    assert len((await list_groups(db, tenant_a.id, limit=500))["groups"]) == 101
