"""Resolved by a person (decided by Aiden, 2026-10-27/28).

When a person fixes an order outside the app, or accepts what is left, they re-check it and accept the
remaining difference. The case is then "resolved": it leaves Needs review and the issue groups, counts as
its own bucket (never as matched), stays resolved while later scans find the same remaining difference,
reopens when that difference changes, and reconciles if the order later matches. Synthetic data only.
"""

from copy import deepcopy
from datetime import timedelta
from uuid import uuid4

import pytest
from sqlalchemy import select, text
from sqlalchemy.exc import DBAPIError

from app.models.audit import AuditEvent
from app.models.transaction_ops import TransactionCase
from app.services.transaction_ops import case_service
from app.services.transaction_ops.state_service import StateError
from tests.test_transaction_cases import NOW, observe, report
from tests.test_transaction_ops_state_db import seed_config

REASON = "fixed_in_netsuite_rest_in_source"


def changed_report(observed):
    body = report(observed=observed)
    body["balance"]["amounts"]["order_total"].update(target="98.00", delta="2.00")
    return body


async def actions(db, tenant, action):
    return list(
        await db.scalars(
            select(AuditEvent).where(AuditEvent.tenant_id == tenant, AuditEvent.action == "transaction_ops." + action)
        )
    )


async def open_case(db, actor):
    config = await seed_config(db, actor.tenant_id, actor)
    await observe(db, actor, config, report(), NOW)
    [case] = await case_service.list_cases(db, actor.tenant_id)
    return config, case


async def resolve(db, actor, case, *, at=None, **overrides):
    values = {
        "reason": REASON,
        "note": "Solidus ops asked to void the duplicate adjustment.",
        "accepted": case_service.remaining_difference(case.latest_report_json),
        **overrides,
    }
    return await case_service.resolve_case(
        db, actor.tenant_id, case.id, actor=actor, now=at or NOW + timedelta(minutes=1), **values
    )


async def test_a_person_accepts_the_remaining_difference_after_a_fresh_check(db, admin_user):
    actor = admin_user[0]
    _, case = await open_case(db, actor)
    accepted = case_service.remaining_difference(case.latest_report_json)
    each = {"source": "100.00", "target": "99.00", "delta": "1.00"}
    assert accepted == {"status": "mismatch", "amounts": {"order_total": each, "tax": each, "refunds": each}}
    resolved = await resolve(db, actor, case)
    assert resolved.status == "resolved"
    record = resolved.resolution_json
    assert (record["reason"], record["resolved_by"], record["accepted"]) == (REASON, str(actor.id), accepted)
    assert record["note"] == "Solidus ops asked to void the duplicate adjustment."
    assert record["observed_at"] == case.last_observed_at.isoformat()
    assert len(await actions(db, actor.tenant_id, "case.resolved_by_person")) == 1


@pytest.mark.parametrize(
    "overrides, at_minutes, code",
    [
        ({}, 16, "recheck_required"),  # the remaining difference must be from a fresh check
        ({"accepted": {"status": "mismatch", "amounts": {"order_total": "2.00"}}}, 1, "remaining_difference_changed"),
        ({"reason": "because"}, 1, "invalid_resolution_reason"),
        ({"note": "x" * 1001}, 1, "invalid_resolution_note"),
    ],
)
async def test_an_acceptance_the_evidence_does_not_back_is_refused(db, admin_user, overrides, at_minutes, code):
    actor = admin_user[0]
    _, case = await open_case(db, actor)
    with pytest.raises(StateError) as exc:
        await resolve(db, actor, case, at=NOW + timedelta(minutes=at_minutes), **overrides)
    assert exc.value.code == code
    assert (await case_service.get_case(db, actor.tenant_id, case.id)).status == "open"


async def test_only_an_open_case_can_be_resolved(db, admin_user):
    actor = admin_user[0]
    config, case = await open_case(db, actor)
    await resolve(db, actor, case)
    with pytest.raises(StateError) as exc:
        await resolve(db, actor, case)
    assert exc.value.code == "case_not_open"
    await observe(db, actor, config, report("matched", NOW + timedelta(minutes=2)), NOW + timedelta(minutes=2))
    reconciled = await case_service.get_case(db, actor.tenant_id, case.id)
    assert reconciled.status == "reconciled"
    with pytest.raises(StateError) as exc:
        await resolve(db, actor, reconciled, at=NOW + timedelta(minutes=3))
    assert exc.value.code == "case_not_open"


async def test_a_later_scan_with_the_same_remaining_difference_keeps_it_resolved(db, admin_user):
    actor = admin_user[0]
    config, case = await open_case(db, actor)
    await resolve(db, actor, case)
    later = NOW + timedelta(hours=1)
    await observe(db, actor, config, report(observed=later), later)
    kept = await case_service.get_case(db, actor.tenant_id, case.id)
    assert kept.status == "resolved" and kept.last_observed_at == later
    assert len(await actions(db, actor.tenant_id, "case.kept_resolved")) == 1


async def test_a_changed_remaining_difference_reopens_it(db, admin_user):
    actor = admin_user[0]
    config, case = await open_case(db, actor)
    await resolve(db, actor, case)
    later = NOW + timedelta(hours=1)
    await observe(db, actor, config, changed_report(later), later)
    reopened = await case_service.get_case(db, actor.tenant_id, case.id)
    assert reopened.status == "open"
    assert reopened.resolution_json["reopened"]["why"] == "remaining_difference_changed"
    assert reopened.resolution_json["reason"] == REASON  # the earlier acceptance stays on record
    assert len(await actions(db, actor.tenant_id, "case.reopened_after_resolution")) == 1


async def test_a_resolved_order_that_later_matches_is_reconciled(db, admin_user):
    actor = admin_user[0]
    config, case = await open_case(db, actor)
    await resolve(db, actor, case)
    later = NOW + timedelta(hours=1)
    await observe(db, actor, config, report("matched", later), later)
    assert (await case_service.get_case(db, actor.tenant_id, case.id)).status == "reconciled"


async def test_a_person_can_reopen_a_resolved_case(db, admin_user):
    actor = admin_user[0]
    _, case = await open_case(db, actor)
    await resolve(db, actor, case)
    reopened = await case_service.reopen_case(db, actor.tenant_id, case.id, actor=actor, now=NOW + timedelta(minutes=2))
    assert reopened.status == "open" and reopened.resolution_json["reopened"]["why"] == "by_person"
    assert len(await actions(db, actor.tenant_id, "case.reopened_by_person")) == 1
    with pytest.raises(StateError) as exc:
        await case_service.reopen_case(db, actor.tenant_id, case.id, actor=actor, now=NOW + timedelta(minutes=3))
    assert exc.value.code == "case_not_resolved"


async def test_the_database_still_refuses_to_move_a_reconciled_case(db, admin_user):
    actor = admin_user[0]
    config, case = await open_case(db, actor)
    await observe(db, actor, config, report("matched", NOW + timedelta(minutes=1)), NOW + timedelta(minutes=1))
    with pytest.raises(DBAPIError, match="reconciled_case_is_final"):
        async with db.begin_nested():
            await db.execute(text("UPDATE transaction_cases SET status='resolved' WHERE id=:id"), {"id": case.id})


async def test_review_results_count_a_resolved_case_apart_and_keep_it_out_of_needs_review(db, admin_user, monkeypatch):
    from app.services.transaction_ops.case_groups import list_groups
    from app.services.transaction_ops.case_service import case_scope
    from app.services.transaction_ops.period_review import review_results
    from app.services.transaction_ops.workspace_results import review_page
    from tests.test_transaction_case_groups import seed
    from tests.test_transaction_review_results import evidence
    from tests.test_transaction_review_slices import review

    actor = admin_user[0]
    _, root = await review(db, actor, monkeypatch)
    await seed(db, actor.tenant_id, 3, scope=case_scope(root))
    cases = list(
        await db.scalars(
            select(TransactionCase)
            .where(TransactionCase.tenant_id == actor.tenant_id)
            .order_by(TransactionCase.order_reference)
        )
    )
    for case in cases:
        row = await evidence(
            db, actor, root, case.order_reference, "difference", root.created_at, source_id=str(case.id)
        )
        body = deepcopy(case.latest_report_json)
        body.update(case_id=str(case.id), source={"record_id": str(case.id)})
        row.report_json = {**row.report_json, **body}
    cases[0].status = "resolved"
    cases[0].resolution_json = {"reason": REASON}
    # The same order number resolved under another subsidiary says nothing about this one.
    twin = uuid4()
    db.add(
        TransactionCase(
            id=twin,
            tenant_id=actor.tenant_id,
            case_key=twin.hex,
            order_reference=cases[1].order_reference,
            scope_json={**case_scope(root), "subsidiary_id": "other"},
            status="resolved",
            resolution_json={"reason": REASON},
            first_observed_at=root.created_at,
            last_observed_at=root.created_at,
            latest_report_json=deepcopy(cases[1].latest_report_json),
        )
    )
    await db.flush()
    results = await review_results(db, actor.tenant_id, root.id, limit=5)
    assert results["summary"] == {"checked": 3, "matched": 0, "resolved": 1, "needs_review": 2, "not_verified": 0}
    page = await review_page(db, actor.tenant_id, [root.id], limit=5, status="needs_review")
    assert sorted(item["order_reference"] for item in page["items"]) == sorted(c.order_reference for c in cases[1:])
    resolved = await review_page(db, actor.tenant_id, [root.id], limit=5, status="resolved")
    assert [item["order_reference"] for item in resolved["items"]] == [cases[0].order_reference]
    assert resolved["items"][0]["resolved"] is True
    groups = await list_groups(db, actor.tenant_id, review_run_ids=[str(root.id)], status="needs_review")
    assert sum(group["case_count"] for group in groups["groups"]) == 2


def test_the_api_offers_exactly_the_service_reasons():
    from typing import get_args

    from app.schemas.transaction_runs import CaseResolution

    assert set(get_args(CaseResolution.model_fields["reason"].annotation)) == case_service.RESOLUTION_REASONS


async def test_a_person_resolves_and_reopens_a_case_over_the_api(app, client, db, admin_user):
    from datetime import datetime, timezone

    from app.api.v1.transaction_ops import router
    from tests.conftest import enable_feature_flag

    app.include_router(router, prefix="/api/v1")
    actor, headers = admin_user
    for flag in ("celigo", "reconciliation"):
        await enable_feature_flag(db, actor.tenant_id, flag)
    config = await seed_config(db, actor.tenant_id, actor)

    fresh = datetime.now(timezone.utc)
    await observe(db, actor, config, report(observed=fresh), fresh)
    [case] = await case_service.list_cases(db, actor.tenant_id)
    case_id, accepted = case.id, case_service.remaining_difference(case.latest_report_json)
    stale = await client.post(
        f"/api/v1/transaction-ops/cases/{case_id}/resolve",
        headers=headers,
        json={"reason": REASON, "accepted": {**accepted, "status": "difference"}},
    )
    assert stale.status_code == 409 and "remaining_difference_changed" in stale.text
    done = await client.post(
        f"/api/v1/transaction-ops/cases/{case_id}/resolve",
        headers=headers,
        json={"reason": REASON, "note": "Solidus ops", "accepted": accepted},
    )
    assert done.status_code == 200, done.text[:300]
    assert done.json()["status"] == "resolved" and done.json()["resolution_json"]["reason"] == REASON
    bad = await client.post(
        f"/api/v1/transaction-ops/cases/{case_id}/resolve",
        headers=headers,
        json={"reason": "because", "accepted": accepted},
    )
    assert bad.status_code == 422
    back = await client.post(f"/api/v1/transaction-ops/cases/{case_id}/reopen", headers=headers)
    assert back.status_code == 200 and back.json()["status"] == "open"


async def test_the_case_offers_exactly_what_may_be_accepted_and_until_when(db, admin_user):
    actor = admin_user[0]
    _, case = await open_case(db, actor)
    offer = case_service.resolution_offer(case, now=NOW + timedelta(minutes=1))
    assert offer == {
        "available": True,
        "unavailable_reason": None,
        "accepted": case_service.remaining_difference(case.latest_report_json),
        "fresh_until": (NOW + case_service.RESOLUTION_FRESHNESS).isoformat(),
        "reasons": sorted(case_service.RESOLUTION_REASONS),
    }
    stale = case_service.resolution_offer(case, now=NOW + timedelta(minutes=16))
    assert (stale["available"], stale["unavailable_reason"]) == (False, "recheck_required")
    resolved = await resolve(db, actor, case)
    assert (
        case_service.resolution_offer(resolved, now=NOW + timedelta(minutes=2))["unavailable_reason"] == "case_not_open"
    )


async def test_reading_a_case_carries_its_offer(app, client, db, admin_user):
    from datetime import datetime, timezone

    from app.api.v1.transaction_ops import router
    from tests.conftest import enable_feature_flag

    app.include_router(router, prefix="/api/v1")
    actor, headers = admin_user
    for flag in ("celigo", "reconciliation"):
        await enable_feature_flag(db, actor.tenant_id, flag)
    config = await seed_config(db, actor.tenant_id, actor)
    fresh = datetime.now(timezone.utc)
    await observe(db, actor, config, report(observed=fresh), fresh)
    [case] = await case_service.list_cases(db, actor.tenant_id)
    response = await client.get(f"/api/v1/transaction-ops/cases/{case.id}", headers=headers)
    assert response.status_code == 200, response.text[:300]
    offer = response.json()["resolution_offer"]
    assert offer["available"] is True and offer["accepted"]["amounts"]["order_total"]["delta"] == "1.00"


# Codex review round 1 (2026-10-28) -------------------------------------------------------------------


@pytest.mark.parametrize("read_minutes", [-60, 10])
async def test_only_a_fresh_read_counts_as_a_fresh_check(db, admin_user, read_minutes):
    # F1: saving an old read now (a saved-evidence review reuses reads) is not a re-check; nor is a read
    # stamped in the future.
    actor = admin_user[0]
    config = await seed_config(db, actor.tenant_id, actor)
    await observe(db, actor, config, report(observed=NOW + timedelta(minutes=read_minutes)), NOW)
    [case] = await case_service.list_cases(db, actor.tenant_id)
    with pytest.raises(StateError) as exc:
        await resolve(db, actor, case)
    assert exc.value.code == "recheck_required"
    offer = case_service.resolution_offer(case, now=NOW + timedelta(minutes=1))
    assert offer["unavailable_reason"] == "recheck_required"


async def test_evidence_that_does_not_say_when_it_was_read_cannot_be_accepted(db, admin_user):
    actor = admin_user[0]
    config = await seed_config(db, actor.tenant_id, actor)
    body = report()
    del body["source"]["observed_at"]
    await observe(db, actor, config, body, NOW)
    [case] = await case_service.list_cases(db, actor.tenant_id)
    with pytest.raises(StateError) as exc:
        await resolve(db, actor, case)
    assert exc.value.code == "recheck_required"


@pytest.mark.parametrize("status, source", [("incomplete", "100.00"), ("mismatch", None)])
async def test_a_difference_that_is_not_measured_cannot_be_accepted(db, admin_user, status, source):
    # F2: incomplete evidence is Not verified, and an unknown source amount gives later scans nothing to
    # compare; neither can be accepted.
    actor = admin_user[0]
    config = await seed_config(db, actor.tenant_id, actor)
    body = report()
    body["balance"]["status"] = status
    body["balance"]["amounts"]["tax"]["source"] = source
    await observe(db, actor, config, body, NOW)
    [case] = await case_service.list_cases(db, actor.tenant_id)
    with pytest.raises(StateError) as exc:
        await resolve(db, actor, case)
    assert exc.value.code == "difference_not_measured"
    offer = case_service.resolution_offer(case, now=NOW + timedelta(minutes=1))
    assert (offer["available"], offer["unavailable_reason"]) == (False, "difference_not_measured")


@pytest.mark.parametrize("status", ["missing_in_netsuite", "ambiguous", "currency_mismatch"])
async def test_only_a_measured_amount_difference_can_be_accepted(db, admin_user, status):
    # Round 2 F4: the reconciler stops before reading amounts for these (order_reconciliation._reconcile),
    # so later scans would have nothing to compare an acceptance with. They are fixed, not accepted.
    actor = admin_user[0]
    config = await seed_config(db, actor.tenant_id, actor)
    body = report()
    body["balance"]["status"] = status
    body["balance"]["amounts"] = {
        key: {"source": None, "target": None, "delta": None} for key in body["balance"]["amounts"]
    }
    await observe(db, actor, config, body, NOW)
    [case] = await case_service.list_cases(db, actor.tenant_id)
    with pytest.raises(StateError) as exc:
        await resolve(db, actor, case, reason="duplicate_or_test_order")
    assert exc.value.code == "difference_not_measured"


async def test_the_same_amounts_written_differently_keep_a_case_resolved(db, admin_user):
    # F3: '1.00' and '1' are the same remaining difference.
    actor = admin_user[0]
    config, case = await open_case(db, actor)
    await resolve(db, actor, case)
    later = NOW + timedelta(hours=1)
    body = report(observed=later)
    for amounts in body["balance"]["amounts"].values():
        amounts.update(source="100", target="99.0", delta="1")
    await observe(db, actor, config, body, later)
    assert (await case_service.get_case(db, actor.tenant_id, case.id)).status == "resolved"
    assert len(await actions(db, actor.tenant_id, "case.kept_resolved")) == 1


async def test_only_a_person_of_the_tenant_with_reconciliation_access_resolves_or_reopens(db, admin_user):
    # F4: the service checks the person itself, as investigate_case does; the route's check is not enough.
    from types import SimpleNamespace

    actor = admin_user[0]
    _, case = await open_case(db, actor)
    stranger = SimpleNamespace(id=uuid4(), tenant_id=uuid4())
    accepted = case_service.remaining_difference(case.latest_report_json)
    with pytest.raises(StateError) as exc:
        await case_service.resolve_case(
            db, actor.tenant_id, case.id, actor=stranger, reason=REASON, note=None, accepted=accepted, now=NOW
        )
    assert exc.value.code == "human_actor_required"
    await resolve(db, actor, case)
    with pytest.raises(StateError) as exc:
        await case_service.reopen_case(db, actor.tenant_id, case.id, actor=stranger, now=NOW + timedelta(minutes=2))
    assert exc.value.code == "human_actor_required"
    resolved = await case_service.get_case(db, actor.tenant_id, case.id)
    assert resolved.resolution_json["resolved_by_name"] == actor.full_name


async def test_resolving_reads_the_case_again_under_its_lock(db, admin_user):
    # F5: another session's newer evidence must win over the copy this session already holds.
    from sqlalchemy import update

    actor = admin_user[0]
    _, case = await open_case(db, actor)
    seen = case_service.remaining_difference(case.latest_report_json)
    await db.execute(
        update(TransactionCase)
        .where(TransactionCase.id == case.id)
        .values(latest_report_json=changed_report(NOW))
        .execution_options(synchronize_session=False)
    )
    with pytest.raises(StateError) as exc:
        await resolve(db, actor, case, accepted=seen)
    assert exc.value.code == "remaining_difference_changed"


@pytest.mark.parametrize("raw", ["difference", "matched"])
async def test_the_export_shows_a_resolved_order_as_resolved_and_never_groups_it(
    app, client, db, admin_user, monkeypatch, raw
):
    # F6, F7: the workbook agrees with the review table (a resolved case is resolved, whatever this
    # review's raw verdict was) and a resolved-only export has no issue groups.
    from io import BytesIO

    from openpyxl import load_workbook

    from app.api.v1.transaction_ops import router
    from app.services.transaction_ops.case_service import case_scope
    from tests.test_transaction_workspace_reports import fixture_rows, linked

    app.include_router(router, prefix="/api/v1")
    actor, headers = admin_user
    first, _, cases = await fixture_rows(db, actor, monkeypatch, count=2)
    identity = uuid4()
    resolved = TransactionCase(
        id=identity,
        tenant_id=actor.tenant_id,
        case_key=identity.hex,
        order_reference="R900000002",
        scope_json=case_scope(first),
        status="resolved",
        resolution_json={"reason": REASON},
        first_observed_at=first.created_at,
        last_observed_at=first.created_at,
        latest_report_json=deepcopy(cases[0].latest_report_json),
    )
    db.add(resolved)
    await db.flush()
    await linked(db, actor, first, resolved, status=raw)
    response = await client.post(
        "/api/v1/transaction-ops/review-export",
        headers=headers,
        json={"review_run_ids": [str(first.id)], "status": "resolved"},
    )
    assert response.status_code == 200, response.text[:300]
    wb = load_workbook(BytesIO(response.content))
    findings = {str(row[0].value): row[4].value for row in wb["Reconciliation"].iter_rows(min_row=2)}
    assert list(findings.values()) == ["resolved"]
    assert not list(wb["Issue groups"].iter_rows(min_row=2, values_only=True))


# Codex review round 2 (2026-10-28) -------------------------------------------------------------------


async def test_a_read_dated_after_now_is_not_fresh_even_beside_a_fresh_one(db, admin_user):
    # R2 F5: the oldest read being fresh does not make a future-dated read acceptable.
    actor = admin_user[0]
    config = await seed_config(db, actor.tenant_id, actor)
    body = report(observed=NOW)
    body["targets"][0]["observed_at"] = (NOW + timedelta(hours=1)).isoformat()
    await observe(db, actor, config, body, NOW)
    [case] = await case_service.list_cases(db, actor.tenant_id)
    with pytest.raises(StateError) as exc:
        await resolve(db, actor, case)
    assert exc.value.code == "recheck_required"


@pytest.mark.parametrize("value", ["sNaN", "NaN", "Infinity", "abc"])
async def test_an_unreadable_accepted_amount_is_a_changed_difference_not_an_error(db, admin_user, value):
    # R2 F6
    actor = admin_user[0]
    _, case = await open_case(db, actor)
    accepted = case_service.remaining_difference(case.latest_report_json)
    accepted["amounts"]["tax"]["source"] = value
    with pytest.raises(StateError) as exc:
        await resolve(db, actor, case, accepted=accepted)
    assert exc.value.code == "remaining_difference_changed"


async def test_a_scan_reads_the_case_again_under_its_lock(db, admin_user):
    # R2 F2: another session resolved the case after this one loaded it; the scan's changed difference
    # must still reopen it.
    from sqlalchemy import update

    actor = admin_user[0]
    config, case = await open_case(db, actor)
    resolution = {"reason": REASON, "accepted": case_service.remaining_difference(case.latest_report_json)}
    await db.execute(
        update(TransactionCase)
        .where(TransactionCase.id == case.id)
        .values(status="resolved", resolution_json=resolution)
        .execution_options(synchronize_session=False)
    )
    later = NOW + timedelta(hours=1)
    await observe(db, actor, config, changed_report(later), later)
    row = (
        await db.execute(text("SELECT status, resolution_json FROM transaction_cases WHERE id = :id"), {"id": case.id})
    ).one()
    assert row.status == "open" and row.resolution_json["reopened"]["why"] == "remaining_difference_changed"


async def test_every_locked_read_of_a_case_sees_the_locked_row(db, admin_user):
    # R2 F2, the mechanism: a locked read refreshes an already-loaded case even when the query forgets to ask.
    from sqlalchemy import update

    actor = admin_user[0]
    _, case = await open_case(db, actor)
    await db.execute(
        update(TransactionCase)
        .where(TransactionCase.id == case.id)
        .values(latest_report_json=changed_report(NOW))
        .execution_options(synchronize_session=False)
    )
    locked = await db.scalar(select(TransactionCase).where(TransactionCase.id == case.id).with_for_update())
    assert locked is case
    assert locked.latest_report_json["balance"]["amounts"]["order_total"]["delta"] == "2.00"
