"""Real scheduler, approval, report, PDF/XLSX and delivery; synthetic provider I/O.

No registry executors, dispatcher, calculations or renderers are replaced. The
BigQuery SDK boundary and Drive storage are synthetic: these checks do not prove
customer source accuracy, Google availability or authorize a production run.
"""

from datetime import datetime, timedelta, timezone
from decimal import Decimal
from io import BytesIO
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from sqlalchemy import event, select

from app.core.encryption import encrypt_credentials
from app.models.audit import AuditEvent
from app.models.job import Job
from app.models.mcp_connector import McpConnector
from app.models.report import Report
from app.services.jobs.recovery import _verify_delivery
from app.services.report import inventory_aging as ia
from app.services.report import report_delivery as delivery
from app.workers.tasks import scheduled_jobs as jobs
from app.workers.tasks.report_auto_refresh import sweep_tenant_reports
from tests.api.test_workflow_review import seed
from tests.jobs.test_delivery_reconciliation import EvidenceDrive
from tests.report.test_inventory_aging import _full_fixture
from tests.report.test_report_delivery import _add_sheets_connector
from tests.report.test_report_pdf import _skip_unless_weasyprint_native_libs


def delivery_fixture():
    payloads, params = _full_fixture()
    return payloads, params


class ArtifactDrive(EvidenceDrive):
    """Retains actual rendered bytes for independent artifact/control checks."""

    def __init__(self):
        super().__init__()
        self.lose_receipt = False
        self.content = {}

    async def upload_new(self, **kwargs):
        result = await super().upload_new(**kwargs)
        self.content[result["file_id"]] = kwargs["content"]
        return result

    async def update_existing(self, **kwargs):
        result = await super().update_existing(**kwargs)
        self.content[result["file_id"]] = kwargs["content"]
        return result


async def setup_report_workflow(
    db, client, user, headers, monkeypatch, *, failure=None, with_delivery=True, budget=None
):
    payloads, params = delivery_fixture()
    if failure == "numeric_mismatch":
        payloads["r_items"][0]["inventory_amount"] += 1
    sources = ia.build_sources(params)
    queries = {s["params"]["query"]: rid for rid, s in sources.items()}
    calls = []

    async def query(**kwargs):
        # Actual dispatcher/connector selection has already run. Refuse any
        # unexpected query or credentials rather than returning generic success.
        assert kwargs["project_id"] == "synthetic-report-project"
        assert kwargs["credentials"] == {"fixture": "scheduled-report"}
        rid = queries[kwargs["query"]]
        calls.append(rid)
        if failure == "missing" and rid == "r_items":
            raise RuntimeError("synthetic source unavailable")
        if failure == "late_error" and rid == "r_prior":
            raise RuntimeError("synthetic lost query receipt")
        rows = payloads[rid]
        columns = list(rows[0])
        return {
            "columns": columns,
            "rows": [[r[c] for c in columns] for r in rows],
            "row_count": len(rows),
            "truncated": failure == "partial" and rid == "r_items",
            "bytes_processed": None if failure == "unknown_usage" else 1024,
            "bytes_billed": 1024,
            "job_id": kwargs["job_id"],
            "cache_hit": False,
        }

    monkeypatch.setattr("app.mcp.tools.bigquery_tools.execute_query", query)
    source = McpConnector(
        tenant_id=user.tenant_id,
        provider="bigquery",
        label="Synthetic inventory source",
        server_url="https://bigquery.googleapis.com",
        auth_type="service_account",
        encrypted_credentials=encrypt_credentials(
            {"service_account_json": {"fixture": "scheduled-report"}, "project_id": "synthetic-report-project"}
        ),
        status="active",
        is_enabled=True,
        metadata_json={"project_id": "synthetic-report-project"},
    )
    db.add(source)
    await _add_sheets_connector(db, user.tenant_id, shared_drive_id="synthetic-test-drive")
    drive = ArtifactDrive()
    monkeypatch.setattr(delivery, "_build_drive_client", lambda *args: drive)
    plan = {
        "steps": [
            {"id": "report", "type": "report.compose", "params": {"playbook_key": "inventory_aging", "params": params}},
            {"id": "deliver", "type": "drive.upload", "params": {"report_step": "report"}},
        ]
    }
    if not with_delivery:
        plan["steps"].pop()
    row = await seed(db, user, plan)
    row.name = "Synthetic inventory delivery"
    row.parameters = {"workflow_review_required": True}
    row.cron_expression = "0 6 * * 1"
    row.budget_json = budget or {"seconds": 120}
    await db.flush()
    review = (await client.post(f"/api/v1/schedules/{row.id}/validate", json={}, headers=headers)).json()
    assert review["ready"], review
    approved = await client.post(
        f"/api/v1/schedules/{row.id}/approve",
        json={"plan_hash": review["plan_hash"], "readiness_hash": review["readiness_hash"]},
        headers=headers,
    )
    assert approved.status_code == 200, approved.text
    row.next_run_at = datetime.now(timezone.utc) - timedelta(minutes=1)
    await db.flush()
    return row, source, drive, calls


async def test_due_workflow_report_is_frozen_at_first_insert_without_pdf(client, db, admin_user, monkeypatch):
    user, headers = admin_user
    await setup_report_workflow(db, client, user, headers, monkeypatch, with_delivery=False)
    tid = user.tenant_id
    inserts = []

    def capture_insert(conn, cursor, statement, parameters, context, executemany):
        compiled = context.compiled
        if compiled and compiled.isinsert and compiled.statement.table.name == "reports":
            inserts.extend(dict(p) for p in context.compiled_parameters)

    bind = db.get_bind()
    event.listen(bind, "after_cursor_execute", capture_insert)
    try:
        assert (await jobs.run_due_jobs(db, tid))["ran"] == 1
    finally:
        event.remove(bind, "after_cursor_execute", capture_insert)
    job = (await db.scalars(select(Job).where(Job.tenant_id == tid))).one()
    assert job.result_summary["reason"] == "done"
    assert len(inserts) == 1
    assert inserts[0]["auto_refresh"] == "off"
    assert inserts[0]["source_run_id"] == job.id
    audit = (
        await db.scalars(select(AuditEvent).where(AuditEvent.tenant_id == tid, AuditEvent.action == "report.compose"))
    ).one()
    assert audit.payload["execution_mode"] == "live" and audit.payload["scheduled_run_id"] == str(job.id)
    assert (await sweep_tenant_reports(db, tid, now=datetime.now(timezone.utc) + timedelta(days=7)))["due"] == 0


async def test_report_rejects_conflicting_run_origins_before_source_execution():
    from app.services.report.playbooks import compose_playbook_report

    # No DB exists: rejection must precede any source or persistence operation.
    with pytest.raises(ValueError, match="both a test and a live"):
        await compose_playbook_report(
            None,
            playbook_key="trial_balance",
            params={"period": "Jun 2026"},
            tenant_id=uuid4(),
            actor_id=None,
            test_run_id=uuid4(),
            scheduled_run_id=uuid4(),
        )


@_skip_unless_weasyprint_native_libs
async def test_approved_due_report_delivers_real_artifacts_once_and_stays_frozen(client, db, admin_user, monkeypatch):
    user, headers = admin_user
    row, _, drive, calls = await setup_report_workflow(db, client, user, headers, monkeypatch)
    tid, sid = user.tenant_id, row.id
    stats = await jobs.run_due_jobs(db, tid)
    assert stats["ran"] == 1 and stats["failed"] == 0, stats
    job = (await db.scalars(select(Job).where(Job.tenant_id == tid, Job.parameters["attempt"].as_integer() == 1))).one()
    report = (await db.scalars(select(Report).where(Report.tenant_id == tid))).one()
    assert job.result_summary["reason"] == "done", job.result_summary
    assert calls == list(ia.RESULT_IDS)
    assert report.auto_refresh == "off" and report.source_run_id == job.id
    assert not report.title.startswith("Test ·")
    assert "Sources" in report.rendered_html and "A-G6" in report.rendered_html
    assert report.version == 1
    model = delivery._inventory_aging_model(report)
    assert Decimal(model["trend"]["Acme"][-1]["value_90p"]) == 21000
    assert Decimal(model["trend"]["Acme"][-1]["pct_90p"]) == Decimal("87.5")
    assert Decimal(model["all_locations"]["aged90_value"]) == 27000
    receipt = job.result_summary["outputs"]["deliver"]
    assert drive.content[receipt["pdf_file_id"]].startswith(b"%PDF-")
    book = load_workbook(BytesIO(drive.content[receipt["xlsx_file_id"]]), data_only=True)
    assert book.sheetnames == ["Summary", "Buckets", "Acme", "Globex", "Initech", "Aged 90+", "Method"]
    summary = list(book["Summary"].values)
    totals = [r for r in summary if r[0] == "Location"]
    assert [(r[1], r[2]) for r in totals] == [
        ("Acme", 24000),
        ("Globex", 10000),
        ("Initech", 8000),
        ("All locations", 42000),
    ]
    assert sum(1 for r in book["Aged 90+"].values if "A-G6" in r) == 1  # beyond top-five display
    assert len(list(book["Acme"].values)) == 10  # header + all nine SKUs
    events = (await db.scalars(select(AuditEvent).where(AuditEvent.job_id == job.id))).all()
    artifact = await _verify_delivery(
        db, tid, sid, job.parameters, row.plan_json["steps"][-1], job.result_summary["step_receipts"], events
    )
    assert artifact["pdf_file_id"] == receipt["pdf_file_id"]
    assert artifact["xlsx_file_id"] == receipt["xlsx_file_id"]
    before = list(drive.calls)
    replay = await jobs.run_schedule_now(db, sid, tenant_id=tid, actor_id=user.id, existing_job_id=job.id)
    assert replay.reason == "done" and drive.calls == before
    assert len((await db.scalars(select(Report).where(Report.tenant_id == tid))).all()) == 1
    assert (await jobs.run_due_jobs(db, tid))["ran"] == 0
    row.paused_at = datetime.now(timezone.utc)
    row.pause_reason = "Synthetic operator stop"
    await db.flush()
    assert (await sweep_tenant_reports(db, tid, now=datetime.now(timezone.utc) + timedelta(days=7)))["due"] == 0
    assert calls == list(ia.RESULT_IDS)
    assert drive.calls.count("upload_new") == 2
    assert job.result_summary["usage"]["seconds"] > 0


@pytest.mark.parametrize(
    "failure", ["missing", "partial", "source_changed", "numeric_mismatch", "late_error", "unknown_usage"]
)
async def test_bad_source_never_reaches_report_delivery(client, db, admin_user, monkeypatch, failure):
    user, headers = admin_user
    row, source, drive, calls = await setup_report_workflow(db, client, user, headers, monkeypatch, failure=failure)
    tid = user.tenant_id
    if failure == "source_changed":
        source.metadata_json = {"project_id": "different-source"}
        await db.flush()
    stats = await jobs.run_due_jobs(db, tid)
    reports = (await db.scalars(select(Report).where(Report.tenant_id == tid))).all()
    assert reports == []
    assert drive.calls == []
    job = (await db.scalars(select(Job).where(Job.tenant_id == tid, Job.parameters["attempt"].as_integer() == 1))).one()
    assert job.result_summary["reason"] in {"error", "blocked"}, stats
    assert job.result_summary.get("detail")
    if failure == "source_changed":
        assert calls == []


@pytest.mark.parametrize("scan_budget,expected_queries,expected_reason", [(4096, 4, "done"), (2048, 2, "budget")])
async def test_report_subqueries_consume_shared_scan_budget(
    client, db, admin_user, monkeypatch, scan_budget, expected_queries, expected_reason
):
    user, headers = admin_user
    row, _, drive, calls = await setup_report_workflow(
        db,
        client,
        user,
        headers,
        monkeypatch,
        with_delivery=False,
        budget={"seconds": 120, "bytes_scanned": scan_budget},
    )
    tid = user.tenant_id
    await jobs.run_due_jobs(db, tid)
    job = (await db.scalars(select(Job).where(Job.tenant_id == tid))).one()
    assert job.result_summary["reason"] == expected_reason
    assert len(calls) == expected_queries
    assert job.result_summary["usage"]["bytes_scanned"] == expected_queries * 1024
    assert job.result_summary["usage"]["query_usage_complete"] is True
    assert job.result_summary["usage"]["usd"] is None
    assert job.result_summary["usage"]["cost_status"] == "unpriced"
    receipts = job.result_summary["report_queries"]
    assert [q["maximum_bytes_billed"] for q in receipts] == [scan_budget - 1024 * n for n in range(expected_queries)]
    assert all(q["state"] == "complete" and q["bytes_billed"] == 1024 for q in receipts)
    assert all(q["source_id"] and len(q["query_sha256"]) == 64 and q["provider_job_id"] for q in receipts)
    assert drive.calls == []


async def test_failed_later_query_preserves_spend_and_fences_retry(client, db, admin_user, monkeypatch):
    user, headers = admin_user
    row, _, drive, calls = await setup_report_workflow(
        db, client, user, headers, monkeypatch, with_delivery=False, failure="late_error"
    )
    tid, sid, uid = user.tenant_id, row.id, user.id
    await jobs.run_due_jobs(db, tid)
    job = (await db.scalars(select(Job).where(Job.tenant_id == tid))).one()
    assert job.result_summary["usage"]["known_bytes_scanned"] == 1024
    assert job.result_summary["usage"]["bytes_scanned"] is None
    assert job.result_summary["usage"]["query_usage_complete"] is False
    assert job.result_summary["verification"] == "uncertain"
    assert [q["state"] for q in job.result_summary["report_queries"]] == ["complete", "pending"]
    before = list(calls)
    outcome = await jobs.run_schedule_now(db, sid, tenant_id=tid, actor_id=uid, existing_job_id=job.id)
    assert outcome.reason == "blocked"
    assert calls == before
    assert drive.calls == []
    assert len((await db.scalars(select(Job).where(Job.tenant_id == tid))).all()) == 1


async def test_report_usd_ceiling_blocks_validation_without_querying(client, db, admin_user, monkeypatch):
    user, headers = admin_user
    row, _, drive, calls = await setup_report_workflow(db, client, user, headers, monkeypatch, with_delivery=False)
    row.budget_json = {"usd": 1, "seconds": 120}
    await db.flush()
    response = await client.post(f"/api/v1/schedules/{row.id}/validate", json={}, headers=headers)
    assert response.status_code == 200
    assert response.json()["ready"] is False
    assert "pricing contract" in str(response.json())
    assert calls == [] and drive.calls == []


async def test_report_query_runtime_refuses_usd_budget_before_spend(db, admin_user):
    from app.services.jobs.registry import StepContext
    from app.services.jobs.report_queries import ReportQueryBudgetError, report_queries

    user, _ = admin_user
    ctx = StepContext(job_id=uuid4(), run_id=uuid4(), tenant_id=user.tenant_id, db=db, budget={"usd": 1})
    with pytest.raises(ReportQueryBudgetError, match="pricing contract"):
        with report_queries(ctx) as usage:
            await usage.begin()
    assert (await db.scalars(select(Job).where(Job.tenant_id == user.tenant_id))).all() == []


async def test_report_query_receipt_cannot_borrow_another_tenants_job(db, admin_user):
    from app.services.jobs.registry import StepContext
    from app.services.jobs.report_queries import ReportQueryUnknownError, report_queries

    user, _ = admin_user
    job = Job(tenant_id=user.tenant_id, job_type="scheduled_job", status="running", parameters={})
    db.add(job)
    await db.flush()
    ctx = StepContext(job_id=uuid4(), run_id=job.id, tenant_id=uuid4(), db=db)
    with pytest.raises(ReportQueryUnknownError, match="Report stopped"):
        with report_queries(ctx) as usage:
            # The tenant guard rejects the borrowed job before recording intent.
            # The context boundary must retain that stop even if a caller catches it.
            with pytest.raises(ReportQueryUnknownError, match="running job"):
                await usage.begin()
            assert usage.stopped is True
            assert usage.pending is None
    await db.refresh(job)
    assert job.status == "running"
    assert not job.result_summary
