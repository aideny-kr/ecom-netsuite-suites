"""A reconciled order stays reconciled (decided by the user, 2026-09-30).

Measured on Framework staging over 14 days: 610 cases reconciled and 66 reopened later. In 63 of
the 66 neither Solidus nor NetSuite had changed; a verified fix had reconciled the order on the
invoice less its credit, and the next scheduled scan compared the sales order total instead. A
later scan now never reopens a reconciled case. When a record really changed after reconciliation
(a new refund, an edited order), the case stays reconciled and the change is recorded on its own.
"""

from copy import deepcopy
from datetime import datetime, timedelta, timezone
from types import SimpleNamespace
from uuid import uuid4

import pytest
from sqlalchemy import func, select, update
from sqlalchemy.exc import DBAPIError

from app.models.audit import AuditEvent
from app.models.transaction_ops import TransactionCase, TransactionCaseObservation
from app.schemas.transaction_runs import ProgressUpdate
from app.services.transaction_ops import case_service
from app.services.transaction_ops import state_service as state
from tests.test_transaction_finding_batch import batch_setup, publish  # noqa: F401  (fixture)


async def events(db, tenant, action):
    return list(
        await db.scalars(
            select(AuditEvent).where(AuditEvent.tenant_id == tenant, AuditEvent.action == "transaction_ops." + action)
        )
    )


async def reconcile_all(db, tenant):
    cases = list(
        await db.scalars(
            select(TransactionCase).where(TransactionCase.tenant_id == tenant).order_by(TransactionCase.order_reference)
        )
    )
    for case in cases:
        case.status = "reconciled"
    await db.flush()
    return cases


def with_refunds(reports):
    """Complete refund reads on both sides, as a normal scan records them."""
    for report in reports:
        report["refund_evidence"] = {
            "source": {"complete": True, "amount": "0", "refund_count": 0, "events": []},
            "target": {"complete": True, "amount": "0", "refund_count": 0, "tax_adjustments": [], "request_links": []},
        }
    return reports


def read_at(report, when):
    """A scan that read the order and its NetSuite records at this time."""
    for part in (report["source"], *report["targets"]):
        part["observed_at"] = when.isoformat()
    return report


async def rearm(db, setup):
    actor, _, run, token, refs, _ = setup
    await state.update_progress(
        db,
        actor.tenant_id,
        run.id,
        ProgressUpdate(progress_json={"pending_refs": refs, "processed": 0}),
        lease_token=token,
    )


async def test_a_later_scan_of_the_same_records_keeps_the_case_reconciled(db, batch_setup):  # noqa: F811
    actor, _, _, _, _, reports = batch_setup
    now = datetime.now(timezone.utc)
    await publish(db, batch_setup, now=now)
    cases = await reconcile_all(db, actor.tenant_id)
    proof = [deepcopy(case.latest_report_json) for case in cases]
    await rearm(db, batch_setup)
    again = deepcopy(reports)
    for report in again:
        report["reason_for_test"] = "the scheduled scan compared the sales order total"
    await publish(db, batch_setup, reports=again, now=now + timedelta(seconds=1))
    assert [case.status for case in cases] == ["reconciled"] * 3
    # The reconciling evidence stays the case's evidence; the scan's reading is kept beside it.
    assert [case.latest_report_json for case in cases] == proof
    observations = await db.scalar(
        select(func.count())
        .select_from(TransactionCaseObservation)
        .where(TransactionCaseObservation.tenant_id == actor.tenant_id)
    )
    assert observations == 6
    assert not await events(db, actor.tenant_id, "case.reopened")
    assert not await events(db, actor.tenant_id, "case.changed_after_reconciliation")
    kept = await events(db, actor.tenant_id, "case.kept_reconciled")
    assert len(kept) == 3 and all(event.payload["observation_id"] for event in kept)


@pytest.mark.parametrize(
    "change",
    [
        "solidus_order_updated",
        "solidus_refunds_changed",
        "netsuite_record_updated",
        "netsuite_records_changed",
        "netsuite_credits_changed",
    ],
)
async def test_a_real_change_after_reconciliation_is_recorded_and_the_case_stays_reconciled(db, batch_setup, change):  # noqa: F811
    actor, _, _, _, _, reports = batch_setup
    with_refunds(reports)
    now = datetime.now(timezone.utc)
    await publish(db, batch_setup, now=now)
    cases = await reconcile_all(db, actor.tenant_id)
    proof = [deepcopy(case.latest_report_json) for case in cases]
    await rearm(db, batch_setup)
    again = deepcopy(reports)
    later = (now + timedelta(days=3)).isoformat()
    for report in again:
        if change == "solidus_order_updated":
            report["source"]["updated_at"] = later
        elif change == "solidus_refunds_changed":
            report["refund_evidence"]["source"]["events"] = [{"id": "30", "amount": "41.38"}]
        elif change == "netsuite_record_updated":
            report["targets"][0]["updated_at"] = later
        elif change == "netsuite_records_changed":
            report["targets"][0]["record_id"] = "999"
        else:
            # A credit memo booked in NetSuite after the order was reconciled.
            report["refund_evidence"]["target"]["tax_adjustments"] = [{"credit_memo_id": "15840539", "amount": "41.38"}]
    await publish(db, batch_setup, reports=again, now=now + timedelta(seconds=1))
    assert [case.status for case in cases] == ["reconciled"] * 3
    assert [case.latest_report_json for case in cases] == proof
    assert not await events(db, actor.tenant_id, "case.reopened")
    assert not await events(db, actor.tenant_id, "case.kept_reconciled")
    flagged = await events(db, actor.tenant_id, "case.changed_after_reconciliation")
    assert len(flagged) == 3
    expected = {"netsuite_record_updated": "netsuite_records_changed"}.get(change, change)
    assert all(event.payload["changes"] == [expected] and event.payload["observation_id"] for event in flagged)
    listed = await case_service.changed_after_reconciliation(db, actor.tenant_id)
    assert sorted(item["case_id"] for item in listed["cases"]) == sorted(str(case.id) for case in cases)
    assert all(item["changes"] == [expected] and item["status"] == "reconciled" for item in listed["cases"])


@pytest.mark.parametrize("unread", ["solidus_refunds", "netsuite_credits", "netsuite_records"])
async def test_evidence_a_scan_could_not_read_is_unknown_not_a_change(db, batch_setup, unread):  # noqa: F811
    # Review finding F1 (#364): a failed refund read or an incomplete lookup is not a change.
    actor, _, _, _, _, reports = batch_setup
    with_refunds(reports)
    now = datetime.now(timezone.utc)
    await publish(db, batch_setup, now=now)
    cases = await reconcile_all(db, actor.tenant_id)
    await rearm(db, batch_setup)
    again = deepcopy(reports)
    for report in again:
        if unread == "solidus_refunds":
            report["refund_evidence"]["source"] = {"complete": False, "reason": "source_refunds_unavailable"}
            report["balance"]["amounts"]["refunds"]["source"] = None
        elif unread == "netsuite_credits":
            report["refund_evidence"]["target"] = {"complete": False, "reason": "target_refunds_unavailable"}
        else:
            report["lookup"] = {**report["lookup"], "complete": False}
            report["targets"] = []
    await publish(db, batch_setup, reports=again, now=now + timedelta(seconds=1))
    assert [case.status for case in cases] == ["reconciled"] * 3
    assert not await events(db, actor.tenant_id, "case.changed_after_reconciliation")
    kept = await events(db, actor.tenant_id, "case.kept_reconciled")
    assert len(kept) == 3 and all(event.payload["unknown"] == [unread] for event in kept)
    assert (await case_service.changed_after_reconciliation(db, actor.tenant_id))["cases"] == []


async def test_a_flag_clears_only_when_a_complete_scan_is_back_at_the_reconciled_records(db, batch_setup):  # noqa: F811
    actor, _, _, _, _, reports = batch_setup
    with_refunds(reports)
    now = datetime.now(timezone.utc)
    await publish(db, batch_setup, reports=[read_at(deepcopy(r), now) for r in reports], now=now)
    await reconcile_all(db, actor.tenant_id)
    steps = []
    for step, edit in enumerate(("changed", "unread", "back")):
        await rearm(db, batch_setup)
        again = [read_at(deepcopy(r), now + timedelta(seconds=step + 1)) for r in reports]
        for report in again:
            report["reason_for_test"] = edit
            if edit == "changed":
                report["source"]["updated_at"] = (now + timedelta(days=3)).isoformat()
            elif edit == "unread":
                report["refund_evidence"]["source"] = {"complete": False, "reason": "source_refunds_unavailable"}
        await publish(db, batch_setup, reports=again, now=now + timedelta(seconds=step + 1))
        steps.append(len((await case_service.changed_after_reconciliation(db, actor.tenant_id))["cases"]))
    # Flagged; an incomplete scan cannot clear it; a complete scan of the reconciled records does.
    assert steps == [3, 3, 0]


async def test_the_single_writer_keeps_a_reconciled_case_reconciled(db, batch_setup):  # noqa: F811
    actor, _, run, token, refs, reports = batch_setup
    now = datetime.now(timezone.utc)
    await state.record_finding(db, actor.tenant_id, run.id, refs[0], reports[0], lease_token=token, now=now)
    case = await db.scalar(
        select(TransactionCase).where(
            TransactionCase.tenant_id == actor.tenant_id, TransactionCase.order_reference == refs[0]
        )
    )
    case.status = "reconciled"
    await db.flush()
    proof = deepcopy(case.latest_report_json)
    changed = deepcopy(reports[0])
    changed["source"]["updated_at"] = (now + timedelta(days=1)).isoformat()
    finding = SimpleNamespace(id=uuid4(), order_reference=refs[0], report_json=changed)
    await case_service.observe_finding(db, actor.tenant_id, run, finding, now=now + timedelta(seconds=1))
    assert case.status == "reconciled" and case.latest_report_json == proof
    assert not await events(db, actor.tenant_id, "case.reopened")
    flagged = await events(db, actor.tenant_id, "case.changed_after_reconciliation")
    assert [event.payload["changes"] for event in flagged] == [["solidus_order_updated"]]


async def test_an_open_case_still_takes_the_newer_evidence(db, batch_setup):  # noqa: F811
    actor, _, _, _, _, reports = batch_setup
    now = datetime.now(timezone.utc)
    await publish(db, batch_setup, now=now)
    await rearm(db, batch_setup)
    again = deepcopy(reports)
    for report in again:
        report["reason_for_test"] = "newer"
    await publish(db, batch_setup, reports=again, now=now + timedelta(seconds=1))
    cases = list(await db.scalars(select(TransactionCase).where(TransactionCase.tenant_id == actor.tenant_id)))
    assert all(case.status == "open" and case.latest_report_json["reason_for_test"] == "newer" for case in cases)
    assert not await events(db, actor.tenant_id, "case.kept_reconciled")


async def test_the_database_refuses_to_reopen_a_reconciled_case(db, batch_setup):  # noqa: F811
    actor = batch_setup[0]
    await publish(db, batch_setup)
    cases = await reconcile_all(db, actor.tenant_id)
    with pytest.raises(DBAPIError, match="reconciled_case_is_final"):
        async with db.begin_nested():
            await db.execute(update(TransactionCase).where(TransactionCase.id == cases[0].id).values(status="open"))


async def test_review_results_and_groups_show_a_reconciled_case_as_matched(db, admin_user, monkeypatch):
    from app.services.transaction_ops.case_groups import list_groups
    from app.services.transaction_ops.case_service import case_scope
    from app.services.transaction_ops.period_review import review_results
    from app.services.transaction_ops.workspace_results import review_page
    from tests.test_transaction_case_groups import seed
    from tests.test_transaction_review_results import evidence
    from tests.test_transaction_review_slices import review

    actor = admin_user[0]
    _, root = await review(db, actor, monkeypatch)
    await seed(db, actor.tenant_id, 3, scope=case_scope(root))
    cases = list(
        await db.scalars(
            select(TransactionCase)
            .where(TransactionCase.tenant_id == actor.tenant_id)
            .order_by(TransactionCase.order_reference)
        )
    )
    for case in cases:
        row = await evidence(
            db, actor, root, case.order_reference, "difference", root.created_at, source_id=str(case.id)
        )
        report = deepcopy(case.latest_report_json)
        report.update(case_id=str(case.id), source={"record_id": str(case.id)})
        row.report_json = {**row.report_json, **report}
    cases[0].status = "reconciled"
    # The same order number reconciled under another subsidiary says nothing about this one.
    twin = uuid4()
    db.add(
        TransactionCase(
            id=twin,
            tenant_id=actor.tenant_id,
            case_key=twin.hex,
            order_reference=cases[1].order_reference,
            scope_json={**case_scope(root), "subsidiary_id": "other"},
            status="reconciled",
            first_observed_at=root.created_at,
            last_observed_at=root.created_at,
            latest_report_json=deepcopy(cases[1].latest_report_json),
        )
    )
    await db.flush()
    results = await review_results(db, actor.tenant_id, root.id, limit=5)
    assert results["summary"] == {"checked": 3, "matched": 1, "needs_review": 2, "not_verified": 0}
    page = await review_page(db, actor.tenant_id, [root.id], limit=5, status="needs_review")
    assert sorted(item["order_reference"] for item in page["items"]) == sorted(c.order_reference for c in cases[1:])
    groups = await list_groups(db, actor.tenant_id, review_run_ids=[str(root.id)], status="needs_review")
    assert sum(group["case_count"] for group in groups["groups"]) == 2


async def test_the_changed_after_reconciliation_list_is_served_to_the_tenant_only(
    app, client, db, admin_user, admin_user_b
):
    from app.api.v1.transaction_ops import router
    from tests.conftest import enable_feature_flag
    from tests.test_transaction_cases import NOW, observe, report
    from tests.test_transaction_ops_state_db import seed_config

    app.include_router(router, prefix="/api/v1")
    actor, headers = admin_user
    other, other_headers = admin_user_b
    for user in (actor, other):
        for feature in ("celigo", "reconciliation"):
            await enable_feature_flag(db, user.tenant_id, feature)
    config = await seed_config(db, actor.tenant_id, actor)
    await observe(db, actor, config, report(), NOW)
    matched = report("matched", NOW + timedelta(seconds=1))
    matched["source"]["updated_at"] = NOW.isoformat()
    await observe(db, actor, config, matched, NOW + timedelta(seconds=1))
    later = report(observed=NOW + timedelta(seconds=2))
    later["source"]["updated_at"] = (NOW + timedelta(days=2)).isoformat()
    await observe(db, actor, config, later, NOW + timedelta(seconds=2))
    path = "/api/v1/transaction-ops/cases/changed-after-reconciliation"
    assert (await client.get(path)).status_code == 401
    response = await client.get(path, headers=headers)
    assert response.status_code == 200, response.text
    [item] = response.json()["cases"]
    assert item["status"] == "reconciled" and item["changes"] == ["solidus_order_updated"]
    assert (await client.get(path, headers=other_headers)).json() == {"cases": []}


async def test_the_export_and_results_page_show_a_reconciled_order_as_reconciled(
    app, client, db, admin_user, monkeypatch
):
    # Review findings F2 and F3 (#364): the workbook and the table label rows from the raw verdict.
    from io import BytesIO

    from openpyxl import load_workbook

    from app.api.v1.transaction_ops import router
    from app.services.transaction_ops.case_service import case_scope
    from app.services.transaction_ops.workspace_results import review_page
    from tests.test_transaction_workspace_reports import fixture_rows, linked

    app.include_router(router, prefix="/api/v1")
    actor, headers = admin_user
    first, second, cases = await fixture_rows(db, actor, monkeypatch, count=2)
    identity = uuid4()
    reconciled = TransactionCase(
        id=identity,
        tenant_id=actor.tenant_id,
        case_key=identity.hex,
        order_reference="R900000001",
        scope_json=case_scope(first),
        status="reconciled",
        first_observed_at=first.created_at,
        last_observed_at=first.created_at,
        latest_report_json=deepcopy(cases[0].latest_report_json),
    )
    db.add(reconciled)
    await db.flush()
    await linked(db, actor, first, reconciled)
    page = await review_page(db, actor.tenant_id, [first.id, second.id], limit=10)
    flags = {item["order_reference"]: item["reconciled"] for item in page["items"]}
    assert flags == {"R900000001": True, cases[0].order_reference: False, cases[1].order_reference: False}
    assert page["summary"]["matched"] == 1 and page["summary"]["needs_review"] == 2
    response = await client.post(
        "/api/v1/transaction-ops/review-export",
        headers=headers,
        json={"review_run_ids": [str(first.id), str(second.id)]},
    )
    assert response.status_code == 200, response.text[:300]
    wb = load_workbook(BytesIO(response.content))
    findings = {str(row[0].value): row[4].value for row in wb["Reconciliation"].iter_rows(min_row=2)}
    [reference] = [key for key in findings if "R900000001" in key]
    assert findings[reference] == "matched"
    assert sum(row[5].value for row in wb["Issue groups"].iter_rows(min_row=2)) == 2


async def test_a_complete_matching_scan_clears_a_flag(db, admin_user):
    # Changed-scope review finding F1 (#364): a clearing read of a reconciled case must clear the flag.
    from tests.test_transaction_cases import NOW, observe, report
    from tests.test_transaction_ops_state_db import seed_config

    actor = admin_user[0]
    config = await seed_config(db, actor.tenant_id, actor)

    def read(status, second, edited):
        body = report(status, NOW + timedelta(seconds=second))
        body["source"]["updated_at"] = edited.isoformat()
        return body

    await observe(db, actor, config, report(), NOW)
    await observe(db, actor, config, read("matched", 1, NOW), NOW + timedelta(seconds=1))
    await observe(db, actor, config, read("mismatch", 2, NOW + timedelta(days=1)), NOW + timedelta(seconds=2))
    assert len((await case_service.changed_after_reconciliation(db, actor.tenant_id))["cases"]) == 1
    await observe(db, actor, config, read("matched", 3, NOW + timedelta(days=1)), NOW + timedelta(seconds=3))
    assert (await case_service.changed_after_reconciliation(db, actor.tenant_id))["cases"] == []


async def test_an_older_read_saved_later_does_not_clear_a_newer_flag(db, batch_setup):  # noqa: F811
    # Changed-scope review finding F4 (#364): the list follows when evidence was read, not when saved.
    actor, _, _, _, _, reports = batch_setup
    with_refunds(reports)
    base = datetime.now(timezone.utc)

    await publish(db, batch_setup, reports=[read_at(deepcopy(r), base) for r in reports], now=base)
    await reconcile_all(db, actor.tenant_id)
    await rearm(db, batch_setup)
    changed = [read_at(deepcopy(r), base + timedelta(seconds=3)) for r in reports]
    for report in changed:
        report["source"]["updated_at"] = (base + timedelta(days=3)).isoformat()
    await publish(db, batch_setup, reports=changed, now=base + timedelta(seconds=3))
    await rearm(db, batch_setup)
    stale = [read_at(deepcopy(r), base + timedelta(seconds=2)) for r in reports]
    for report in stale:
        report["reason_for_test"] = "an older read that finished late"
    await publish(db, batch_setup, reports=stale, now=base + timedelta(seconds=4))
    listed = await case_service.changed_after_reconciliation(db, actor.tenant_id)
    assert len(listed["cases"]) == 3
    assert all(item["changes"] == ["solidus_order_updated"] for item in listed["cases"])


@pytest.mark.parametrize(
    "variant, changes, unknown",
    [
        # Round-3 packet review of #364 (F5-F8), and the only difference seen on real unchanged reads.
        ("refund_event_details_unread", [], ["solidus_refunds"]),
        ("netsuite_edit_time_missing", [], ["netsuite_records"]),
        ("netsuite_order_vanished", ["netsuite_records_changed"], []),
        ("netsuite_refund_records_replaced", ["netsuite_credits_changed"], []),
        ("read_metadata_only", [], []),
    ],
)
def test_changes_since_compares_everything_that_was_read(variant, changes, unknown):
    from tests.test_transaction_ops_planner import planning_case

    before = with_refunds([deepcopy(planning_case().report)])[0]
    before["refund_evidence"]["source"].update(events=[{"id": "30", "amount": "20"}], events_complete=True)
    before["refund_evidence"]["target"].update(record_ids=["4"], api_calls=3)
    after = deepcopy(before)
    if variant == "refund_event_details_unread":
        after["refund_evidence"]["source"].update(events=[], events_complete=False)
    elif variant == "netsuite_edit_time_missing":
        after["targets"][0]["updated_at"] = None
    elif variant == "netsuite_order_vanished":
        after["targets"] = []
        after["lookup"] = {**after["lookup"], "count": 0}
    elif variant == "netsuite_refund_records_replaced":
        after["refund_evidence"]["target"]["record_ids"] = ["5"]
    else:
        after["refund_evidence"]["target"].update(api_calls=7, dependency_manifest={"refund_requests": ["20"]})
        after["source"]["observed_at"] = after["targets"][0]["observed_at"] = datetime.now(timezone.utc).isoformat()
    assert case_service.changes_since(before, after) == (changes, unknown)
