"""Task 5 (Slice 1) — Drive delivery service.

Spec: docs/superpowers/specs/2026-09-08-scheduled-jobs-and-inventory-aging-design.md
§A5. ``deliver_report_to_drive(db, *, tenant_id, report_id, actor_type, actor_id,
period_key) -> DeliveryResult`` uploads a PDF + Excel workbook of a composed report to
the tenant's Google Drive, idempotently keyed by ``period_key`` (a same-named existing
file is UPDATED, never duplicated), with a `report.delivery.started` audit event
written BEFORE any Drive call and `report.delivery.completed`/`failed` after.

``deliver_report_to_drive`` reaches Google Drive through the module-level
``DriveClient`` protocol (find_folder/create_folder/find_file/upload_new/
update_existing) via the patchable factory ``report_delivery._build_drive_client`` —
these tests monkeypatch that factory to return ``FakeDriveClient`` below, an in-memory
stand-in that also records call order into a shared list alongside audit-event writes,
so ordering assertions don't need real Google API calls. The PDF/Excel BYTE renderers
(``report_delivery._render_pdf_bytes`` / ``_render_xlsx_bytes``) are patched too: real
PDF rendering goes through WeasyPrint (Task 4), which needs native pango/cairo/
gdk-pixbuf libraries this macOS dev machine does not have (see
``tests/report/test_report_pdf.py``'s own skip-probe) — Task 5 owns the Drive-delivery
MECHANICS (folder/file idempotency, audit ordering, failure handling), not report
rendering, so these tests stub the byte producers rather than skip wholesale.
"""

from __future__ import annotations

from datetime import datetime

import pytest
from sqlalchemy import select, text

from app.core.database import set_tenant_context
from app.core.encryption import encrypt_credentials
from app.models.mcp_connector import McpConnector
from app.models.report import Report
from app.services.report import report_delivery
from tests.conftest import create_test_tenant, create_test_user


async def _seed_report(db, tenant, user, *, title: str = "Inventory Aging Weekly") -> Report:
    report = Report(
        tenant_id=tenant.id,
        title=title,
        spec_json={"title": title, "sections": []},
        rendered_html="<html><body>REPORT</body></html>",
        created_by=user.id,
    )
    db.add(report)
    await db.flush()
    return report


async def _seed_inventory_aging_report(db, tenant, user) -> Report:
    """A real inventory_aging report row -- spec_json/recipe_json built exactly the
    way compose_playbook_report persists them (compute() -> build_inventory_aging_sections
    -> spec_json_safe), without going through the tool-dispatch machinery Task 5's
    delivery service doesn't own. Gate fix #6's own test needs this (not the plain
    _seed_report above, whose empty sections carry no playbook recipe at all)."""
    from app.services.report.inventory_aging import compute
    from app.services.report.report_html import build_inventory_aging_sections
    from app.services.report.report_service import spec_json_safe
    from tests.report.test_inventory_aging import _full_fixture

    payloads, params = _full_fixture()
    report_data = compute(payloads, params)
    sections = build_inventory_aging_sections(report_data, composed_at="2026-09-08T13:00:00+00:00")
    safe_spec = spec_json_safe({"title": "Inventory Aging Weekly", "sections": sections})
    report = Report(
        tenant_id=tenant.id,
        title="Inventory Aging Weekly",
        spec_json=safe_spec,
        rendered_html="<html><body>REPORT</body></html>",
        created_by=user.id,
        recipe_json={
            "schema_version": 1,
            "captured_at": "2026-09-08T13:00:00+00:00",
            "playbook": {"key": "inventory_aging", "params": params},
            "sections": [{"type": "watch_items", "result_ids": ["r_items"], "params": params}],
            "sources": {"r_items": {"tool": "bigquery_sql", "params": {"query": "SELECT 1"}, "connection_id": None}},
        },
    )
    db.add(report)
    await db.flush()
    return report


async def _add_sheets_connector(db, tenant_id, *, shared_drive_id: str | None = None) -> McpConnector:
    encrypted = encrypt_credentials({"service_account_json": {"client_email": "sa@test.iam.gserviceaccount.com"}})
    connector = McpConnector(
        tenant_id=tenant_id,
        provider="google_sheets",
        label="test drive",
        server_url="https://sheets.googleapis.com",
        auth_type="service_account",
        encrypted_credentials=encrypted,
        encryption_key_version=1,
        status="active",
        is_enabled=True,
        metadata_json={"shared_drive_id": shared_drive_id} if shared_drive_id else {},
    )
    db.add(connector)
    await db.flush()
    return connector


class FakeDriveClient:
    """In-memory DriveClient. Gate fix #5: folders/files are keyed by their
    ``app_properties`` identity (report_series_id/report_id for a folder;
    report_id/period_key/kind for a file) when given — the SAME identity scheme
    ``deliver_report_to_drive`` now keys on, so this fake reproduces the real
    collision fix rather than papering over it. ``app_properties=None`` (the
    singleton top-level "Reports" folder) falls back to name-based keying,
    unchanged. Records every call into ``calls`` (shared with the test) so
    ordering vs. the audit-event write can be asserted without a real Google API
    round-trip."""

    def __init__(self, calls: list[str], *, fail_on: str | None = None):
        self.calls = calls
        self.fail_on = fail_on
        self._folders: dict[tuple, str] = {}
        self._files: dict[tuple, dict[str, str]] = {}
        self._next_id = 0

    def _new_id(self, prefix: str) -> str:
        self._next_id += 1
        return f"{prefix}-{self._next_id}"

    def _maybe_fail(self, call: str):
        if self.fail_on == call:
            raise RuntimeError(f"simulated Drive failure at {call}")

    @staticmethod
    def _key(name: str, parent_id: str | None, app_properties: dict[str, str] | None) -> tuple:
        identity = tuple(sorted(app_properties.items())) if app_properties else (("name", name),)
        return (identity, parent_id)

    async def find_folder(
        self, *, name: str, parent_id: str | None, app_properties: dict[str, str] | None = None
    ) -> str | None:
        self.calls.append("find_folder")
        self._maybe_fail("find_folder")
        return self._folders.get(self._key(name, parent_id, app_properties))

    async def create_folder(
        self, *, name: str, parent_id: str | None, app_properties: dict[str, str] | None = None
    ) -> str:
        self.calls.append("create_folder")
        self._maybe_fail("create_folder")
        folder_id = self._new_id("folder")
        self._folders[self._key(name, parent_id, app_properties)] = folder_id
        return folder_id

    async def find_file(
        self, *, name: str, parent_id: str, app_properties: dict[str, str] | None = None
    ) -> dict[str, str] | None:
        self.calls.append("find_file")
        self._maybe_fail("find_file")
        return self._files.get(self._key(name, parent_id, app_properties))

    async def upload_new(
        self,
        *,
        name: str,
        parent_id: str,
        content: bytes,
        mime_type: str,
        app_properties: dict[str, str] | None = None,
    ) -> dict[str, str]:
        self.calls.append("upload_new")
        self._maybe_fail("upload_new")
        file_id = self._new_id("file")
        record = {"file_id": file_id, "url": f"https://drive.example/{file_id}"}
        self._files[self._key(name, parent_id, app_properties)] = record
        return record

    async def update_existing(self, *, file_id: str, content: bytes, mime_type: str) -> dict[str, str]:
        self.calls.append("update_existing")
        self._maybe_fail("update_existing")
        for record in self._files.values():
            if record["file_id"] == file_id:
                return record
        return {"file_id": file_id, "url": f"https://drive.example/{file_id}"}


def _patch_renderers(monkeypatch, calls: list[str]):
    def fake_pdf(report) -> bytes:
        calls.append("render_pdf")
        return b"%PDF-FAKE"

    def fake_xlsx(report) -> bytes:
        calls.append("render_xlsx")
        return b"XLSX-FAKE"

    monkeypatch.setattr(report_delivery, "_render_pdf_bytes", fake_pdf)
    monkeypatch.setattr(report_delivery, "_render_xlsx_bytes", fake_xlsx)


def _patch_drive_client(monkeypatch, client: FakeDriveClient):
    def factory(credentials, shared_drive_id):
        return client

    monkeypatch.setattr(report_delivery, "_build_drive_client", factory)


async def test_no_connector_raises_delivery_unavailable(db, monkeypatch):
    tenant = await create_test_tenant(db, name="NoConnector")
    user, _ = await create_test_user(db, tenant)
    await set_tenant_context(db, str(tenant.id))
    report = await _seed_report(db, tenant, user)

    calls: list[str] = []
    _patch_renderers(monkeypatch, calls)

    with pytest.raises(report_delivery.DeliveryUnavailable):
        await report_delivery.deliver_report_to_drive(
            db,
            tenant_id=tenant.id,
            report_id=report.id,
            actor_type="user",
            actor_id=user.id,
            period_key="2026-09-08",
        )
    # nothing rendered, nothing recorded — the check fails before any work starts
    assert calls == []


async def test_folder_found_or_created_only_once_across_two_deliveries(db, monkeypatch):
    tenant = await create_test_tenant(db, name="FolderOnce")
    user, _ = await create_test_user(db, tenant)
    await set_tenant_context(db, str(tenant.id))
    report = await _seed_report(db, tenant, user)
    await _add_sheets_connector(db, tenant.id)

    calls: list[str] = []
    _patch_renderers(monkeypatch, calls)
    client = FakeDriveClient(calls)
    _patch_drive_client(monkeypatch, client)

    await report_delivery.deliver_report_to_drive(
        db, tenant_id=tenant.id, report_id=report.id, actor_type="user", actor_id=user.id, period_key="2026-09-08"
    )
    first_creates = calls.count("create_folder")
    assert first_creates == 2  # "Reports" + the report's own title, both new

    await set_tenant_context(db, str(tenant.id))
    calls.clear()
    await report_delivery.deliver_report_to_drive(
        db, tenant_id=tenant.id, report_id=report.id, actor_type="user", actor_id=user.id, period_key="2026-09-08"
    )
    assert calls.count("create_folder") == 0
    assert calls.count("find_folder") == 2


async def test_two_reports_with_the_same_title_get_separate_folders_and_files(db, monkeypatch):
    """Gate fix #5: folder/file lookup used to key on report.title + period_key alone
    -- two DIFFERENT reports sharing a title (e.g. two inventory_aging series, or a
    renamed/recomposed report reusing a common name) collided on the SAME Drive
    folder and files, silently overwriting one report's delivery with the other's.
    Identity must be the report (series_id, or report.id when there is no series),
    not the human-readable title."""
    tenant = await create_test_tenant(db, name="SameTitleCorp")
    user, _ = await create_test_user(db, tenant)
    await set_tenant_context(db, str(tenant.id))
    report_a = await _seed_report(db, tenant, user, title="Inventory Aging Weekly")
    report_b = await _seed_report(db, tenant, user, title="Inventory Aging Weekly")
    await _add_sheets_connector(db, tenant.id)

    calls: list[str] = []
    _patch_renderers(monkeypatch, calls)
    client = FakeDriveClient(calls)
    _patch_drive_client(monkeypatch, client)

    result_a = await report_delivery.deliver_report_to_drive(
        db, tenant_id=tenant.id, report_id=report_a.id, actor_type="user", actor_id=user.id, period_key="2026-09-08"
    )
    await set_tenant_context(db, str(tenant.id))
    result_b = await report_delivery.deliver_report_to_drive(
        db, tenant_id=tenant.id, report_id=report_b.id, actor_type="user", actor_id=user.id, period_key="2026-09-08"
    )

    assert result_a.folder_id != result_b.folder_id
    assert result_a.pdf_file_id != result_b.pdf_file_id
    assert result_a.xlsx_file_id != result_b.xlsx_file_id


async def test_redelivery_of_the_same_report_and_period_still_updates_in_place(db, monkeypatch):
    """Companion to the collision fix above: identity-by-report must not break the
    existing idempotency contract -- the SAME report + period re-delivered still
    resolves to the SAME folder/files (update, never a duplicate)."""
    tenant = await create_test_tenant(db, name="RedeliverSameCorp")
    user, _ = await create_test_user(db, tenant)
    await set_tenant_context(db, str(tenant.id))
    report = await _seed_report(db, tenant, user, title="Inventory Aging Weekly")
    await _add_sheets_connector(db, tenant.id)

    calls: list[str] = []
    _patch_renderers(monkeypatch, calls)
    client = FakeDriveClient(calls)
    _patch_drive_client(monkeypatch, client)

    first = await report_delivery.deliver_report_to_drive(
        db, tenant_id=tenant.id, report_id=report.id, actor_type="user", actor_id=user.id, period_key="2026-09-08"
    )
    await set_tenant_context(db, str(tenant.id))
    second = await report_delivery.deliver_report_to_drive(
        db, tenant_id=tenant.id, report_id=report.id, actor_type="user", actor_id=user.id, period_key="2026-09-08"
    )

    assert second.folder_id == first.folder_id
    assert second.pdf_file_id == first.pdf_file_id
    assert second.xlsx_file_id == first.xlsx_file_id


async def test_deliver_report_to_drive_takes_a_per_report_advisory_lock_before_any_drive_call(db, monkeypatch):
    """Gate fix #7: two concurrent deliveries of the SAME report can each run
    _upload_or_update's find-then-create sequence for the folder and both files —
    without serialization, both could see "nothing yet" and both upload, duplicating
    every Drive artifact instead of one updating in place (test_redelivery_... above
    proves find-then-create itself is correct once serialized; this proves the
    serialization actually happens). deliver_report_to_drive must take a per-report
    Postgres advisory TRANSACTION lock (pg_advisory_xact_lock(hashtext(report_id))),
    acquired before the folder lookup and released only at the delivery's own
    commit/rollback (xact-scoped), so a genuinely concurrent second caller's Drive
    work cannot even START until the first's has fully landed.

    A real cross-connection block can't be demonstrated with this repo's `db`
    fixture (conftest.py wraps every test in one savepoint-scoped connection that
    never truly commits, and a Postgres advisory xact lock is session-scoped —
    re-acquiring it on the SAME session never blocks itself) — this instead proves
    the mechanism is wired in correctly: the lock statement runs, with a key
    derived from this report's id, and BEFORE any find/create/upload call."""
    tenant = await create_test_tenant(db, name="AdvisoryLockCorp")
    user, _ = await create_test_user(db, tenant)
    await set_tenant_context(db, str(tenant.id))
    report = await _seed_report(db, tenant, user)
    await _add_sheets_connector(db, tenant.id)

    calls: list[str] = []
    _patch_renderers(monkeypatch, calls)
    client = FakeDriveClient(calls)
    _patch_drive_client(monkeypatch, client)

    real_execute = db.execute

    async def spy_execute(stmt, *args, **kwargs):
        sql_text = str(getattr(stmt, "text", stmt))
        if "pg_advisory_xact_lock" in sql_text:
            calls.append("advisory_lock")
            params = args[0] if args else kwargs.get("parameters") or {}
            assert params.get("report_id") == str(report.id)
        return await real_execute(stmt, *args, **kwargs)

    monkeypatch.setattr(db, "execute", spy_execute)

    await report_delivery.deliver_report_to_drive(
        db, tenant_id=tenant.id, report_id=report.id, actor_type="user", actor_id=user.id, period_key="2026-09-08"
    )

    assert calls.count("advisory_lock") == 1
    lock_idx = calls.index("advisory_lock")
    drive_idxs = [
        i
        for i, c in enumerate(calls)
        if c in ("find_folder", "create_folder", "find_file", "upload_new", "update_existing")
    ]
    assert drive_idxs, "expected at least one Drive call"
    assert lock_idx < min(drive_idxs)


async def test_inventory_aging_delivery_builds_the_real_seven_sheet_workbook(db, monkeypatch):
    """Gate fix #6: _render_xlsx_bytes built the generic 5-row metadata sheet for
    EVERY report type, including inventory_aging. For an inventory_aging report,
    delivery must build the real seven-sheet workbook
    (report_excel.build_inventory_aging_workbook) from the report's stored
    JSON-safe model (extended in gate fix #4/#6 to accept that form directly)."""
    tenant = await create_test_tenant(db, name="RealWorkbookCorp")
    user, _ = await create_test_user(db, tenant)
    await set_tenant_context(db, str(tenant.id))
    report = await _seed_inventory_aging_report(db, tenant, user)
    await _add_sheets_connector(db, tenant.id)

    calls: list[str] = []

    def fake_pdf(report) -> bytes:
        calls.append("render_pdf")
        return b"%PDF-FAKE"

    monkeypatch.setattr(report_delivery, "_render_pdf_bytes", fake_pdf)

    seen_xlsx_bytes: dict[str, bytes] = {}

    class CapturingDriveClient(FakeDriveClient):
        async def upload_new(self, *, name, parent_id, content, mime_type, app_properties=None):
            if mime_type == report_delivery._XLSX_MIME:
                seen_xlsx_bytes["content"] = content
            return await super().upload_new(
                name=name, parent_id=parent_id, content=content, mime_type=mime_type, app_properties=app_properties
            )

    client = CapturingDriveClient(calls)
    _patch_drive_client(monkeypatch, client)

    await report_delivery.deliver_report_to_drive(
        db, tenant_id=tenant.id, report_id=report.id, actor_type="user", actor_id=user.id, period_key="2026-09-08"
    )

    from io import BytesIO

    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(seen_xlsx_bytes["content"]))
    assert len(wb.sheetnames) == 7
    assert wb.sheetnames[0] == "Summary"
    assert wb.sheetnames[1] == "Buckets"
    assert wb.sheetnames[-1] == "Method"


async def test_non_playbook_report_delivery_still_uses_the_generic_workbook(db, monkeypatch):
    """Companion to the fix above: a report with no inventory_aging playbook recipe
    (a chat-composed table/narrative report) must keep the generic single-sheet
    metadata workbook — the type-aware routing must not misfire on every report."""
    tenant = await create_test_tenant(db, name="GenericWorkbookCorp")
    user, _ = await create_test_user(db, tenant)
    await set_tenant_context(db, str(tenant.id))
    report = await _seed_report(db, tenant, user, title="Cash report")
    await _add_sheets_connector(db, tenant.id)

    calls: list[str] = []
    _patch_renderers(monkeypatch, calls)

    seen_xlsx_bytes: dict[str, bytes] = {}

    class CapturingDriveClient(FakeDriveClient):
        async def upload_new(self, *, name, parent_id, content, mime_type, app_properties=None):
            if mime_type == report_delivery._XLSX_MIME:
                seen_xlsx_bytes["content"] = content
            return await super().upload_new(
                name=name, parent_id=parent_id, content=content, mime_type=mime_type, app_properties=app_properties
            )

    client = CapturingDriveClient(calls)
    _patch_drive_client(monkeypatch, client)

    await report_delivery.deliver_report_to_drive(
        db, tenant_id=tenant.id, report_id=report.id, actor_type="user", actor_id=user.id, period_key="2026-09-08"
    )

    # _render_xlsx_bytes was patched to the fake byte producer (calls it, doesn't
    # build a real workbook) -- proves the inventory_aging-specific path was never
    # taken for a report with no playbook recipe.
    assert "render_xlsx" in calls
    assert seen_xlsx_bytes["content"] == b"XLSX-FAKE"


async def test_first_delivery_uploads_two_files(db, monkeypatch):
    tenant = await create_test_tenant(db, name="FirstDelivery")
    user, _ = await create_test_user(db, tenant)
    await set_tenant_context(db, str(tenant.id))
    report = await _seed_report(db, tenant, user)
    await _add_sheets_connector(db, tenant.id)

    calls: list[str] = []
    _patch_renderers(monkeypatch, calls)
    client = FakeDriveClient(calls)
    _patch_drive_client(monkeypatch, client)

    result = await report_delivery.deliver_report_to_drive(
        db, tenant_id=tenant.id, report_id=report.id, actor_type="user", actor_id=user.id, period_key="2026-09-08"
    )
    assert calls.count("upload_new") == 2
    assert calls.count("update_existing") == 0
    assert result.pdf_file_id and result.xlsx_file_id
    assert result.pdf_file_id != result.xlsx_file_id
    assert result.pdf_url and result.xlsx_url
    assert result.folder_id
    assert isinstance(result.delivered_at, datetime)


async def test_second_delivery_updates_both_no_new_file_ids(db, monkeypatch):
    tenant = await create_test_tenant(db, name="SecondDelivery")
    user, _ = await create_test_user(db, tenant)
    await set_tenant_context(db, str(tenant.id))
    report = await _seed_report(db, tenant, user)
    await _add_sheets_connector(db, tenant.id)

    calls: list[str] = []
    _patch_renderers(monkeypatch, calls)
    client = FakeDriveClient(calls)
    _patch_drive_client(monkeypatch, client)

    first = await report_delivery.deliver_report_to_drive(
        db, tenant_id=tenant.id, report_id=report.id, actor_type="user", actor_id=user.id, period_key="2026-09-08"
    )

    await set_tenant_context(db, str(tenant.id))
    calls.clear()
    second = await report_delivery.deliver_report_to_drive(
        db, tenant_id=tenant.id, report_id=report.id, actor_type="user", actor_id=user.id, period_key="2026-09-08"
    )

    assert calls.count("upload_new") == 0
    assert calls.count("update_existing") == 2
    assert second.pdf_file_id == first.pdf_file_id
    assert second.xlsx_file_id == first.xlsx_file_id


async def test_started_audit_event_recorded_before_any_upload_call(db, monkeypatch):
    tenant = await create_test_tenant(db, name="AuditOrder")
    user, _ = await create_test_user(db, tenant)
    await set_tenant_context(db, str(tenant.id))
    report = await _seed_report(db, tenant, user)
    await _add_sheets_connector(db, tenant.id)

    calls: list[str] = []
    _patch_renderers(monkeypatch, calls)
    client = FakeDriveClient(calls)
    _patch_drive_client(monkeypatch, client)

    from app.services import audit_service as audit_service_module

    real_log_event = audit_service_module.log_event

    async def wrapped_log_event(*args, **kwargs):
        if kwargs.get("action") == "report.delivery.started":
            calls.append("audit_started")
        return await real_log_event(*args, **kwargs)

    monkeypatch.setattr(report_delivery.audit_service, "log_event", wrapped_log_event)

    await report_delivery.deliver_report_to_drive(
        db, tenant_id=tenant.id, report_id=report.id, actor_type="user", actor_id=user.id, period_key="2026-09-08"
    )

    started_idx = calls.index("audit_started")
    upload_indexes = [i for i, c in enumerate(calls) if c in ("find_folder", "create_folder", "upload_new")]
    assert upload_indexes, "expected at least one Drive call"
    assert started_idx < min(upload_indexes)

    await set_tenant_context(db, str(tenant.id))
    row = (
        await db.execute(
            text(
                "SELECT actor_id, actor_type, payload FROM audit_events "
                "WHERE action='report.delivery.started' AND resource_id=:rid"
            ),
            {"rid": str(report.id)},
        )
    ).first()
    assert row is not None
    assert row[0] == user.id and row[1] == "user"
    assert "idempotency_key" in row[2]


async def test_success_writes_delivery_json_and_published_fields(db, monkeypatch):
    tenant = await create_test_tenant(db, name="SuccessWrite")
    user, _ = await create_test_user(db, tenant)
    await set_tenant_context(db, str(tenant.id))
    report = await _seed_report(db, tenant, user)
    await _add_sheets_connector(db, tenant.id)

    calls: list[str] = []
    _patch_renderers(monkeypatch, calls)
    client = FakeDriveClient(calls)
    _patch_drive_client(monkeypatch, client)

    result = await report_delivery.deliver_report_to_drive(
        db, tenant_id=tenant.id, report_id=report.id, actor_type="user", actor_id=user.id, period_key="2026-09-08"
    )

    await set_tenant_context(db, str(tenant.id))
    row = (await db.execute(select(Report).where(Report.id == report.id))).scalar_one()
    assert row.published_drive_url == result.pdf_url
    assert row.published_at is not None
    assert row.delivery_json is not None
    assert row.delivery_json["pdf"] == {"file_id": result.pdf_file_id, "url": result.pdf_url}
    assert row.delivery_json["xlsx"] == {"file_id": result.xlsx_file_id, "url": result.xlsx_url}
    assert row.delivery_json["folder_id"] == result.folder_id
    assert row.delivery_json["period_key"] == "2026-09-08"
    assert row.delivery_json["delivered_at"]

    audit_row = (
        await db.execute(
            text("SELECT 1 FROM audit_events WHERE action='report.delivery.completed' AND resource_id=:rid"),
            {"rid": str(report.id)},
        )
    ).first()
    assert audit_row is not None


async def test_failing_upload_writes_failed_audit_and_no_partial_delivery_json(db, monkeypatch):
    tenant = await create_test_tenant(db, name="FailingUpload")
    user, _ = await create_test_user(db, tenant)
    await set_tenant_context(db, str(tenant.id))
    report = await _seed_report(db, tenant, user)
    await _add_sheets_connector(db, tenant.id)
    # Capture raw ids BEFORE the call: the service's internal failure-path rollback
    # (SQLAlchemy's default expire_on_rollback=True) expires every ORM object this
    # session tracked, and a plain attribute access on an expired AsyncSession object
    # afterward raises MissingGreenlet (the async ORM's lazy-load bridge needs an
    # explicit await, e.g. db.refresh() — a bare `tenant.id` does not provide one).
    # A real caller never hits this: it evaluates report_id=row.id at the call site,
    # not after the exception, exactly like this test now does.
    tenant_id, report_id, user_id = tenant.id, report.id, user.id

    calls: list[str] = []
    _patch_renderers(monkeypatch, calls)
    client = FakeDriveClient(calls, fail_on="upload_new")
    _patch_drive_client(monkeypatch, client)

    with pytest.raises(report_delivery.DeliveryFailed):
        await report_delivery.deliver_report_to_drive(
            db,
            tenant_id=tenant_id,
            report_id=report_id,
            actor_type="user",
            actor_id=user_id,
            period_key="2026-09-08",
        )

    await set_tenant_context(db, str(tenant_id))
    row = (await db.execute(select(Report).where(Report.id == report_id))).scalar_one()
    assert row.delivery_json is None
    assert row.published_drive_url is None
    assert row.published_at is None

    failed_audit = (
        await db.execute(
            text(
                "SELECT status, error_message FROM audit_events WHERE action='report.delivery.failed' AND resource_id=:rid"
            ),
            {"rid": str(report_id)},
        )
    ).first()
    assert failed_audit is not None
    assert failed_audit[0] == "error"
    assert "simulated Drive failure" in failed_audit[1]

    started_audit = (
        await db.execute(
            text("SELECT 1 FROM audit_events WHERE action='report.delivery.started' AND resource_id=:rid"),
            {"rid": str(report_id)},
        )
    ).first()
    assert started_audit is not None


async def test_shared_drive_id_passed_from_connector_metadata(db, monkeypatch):
    tenant = await create_test_tenant(db, name="SharedDrive")
    user, _ = await create_test_user(db, tenant)
    await set_tenant_context(db, str(tenant.id))
    report = await _seed_report(db, tenant, user)
    await _add_sheets_connector(db, tenant.id, shared_drive_id="0AbCdEfGhIjKlMnOpQrS")

    calls: list[str] = []
    _patch_renderers(monkeypatch, calls)
    client = FakeDriveClient(calls)

    seen = {}

    def factory(credentials, shared_drive_id):
        seen["shared_drive_id"] = shared_drive_id
        return client

    monkeypatch.setattr(report_delivery, "_build_drive_client", factory)

    await report_delivery.deliver_report_to_drive(
        db, tenant_id=tenant.id, report_id=report.id, actor_type="user", actor_id=user.id, period_key="2026-09-08"
    )
    assert seen["shared_drive_id"] == "0AbCdEfGhIjKlMnOpQrS"


async def test_filenames_include_report_title_and_period_key(db, monkeypatch):
    tenant = await create_test_tenant(db, name="Filenames")
    user, _ = await create_test_user(db, tenant)
    await set_tenant_context(db, str(tenant.id))
    report = await _seed_report(db, tenant, user, title="Inventory Aging Weekly")
    await _add_sheets_connector(db, tenant.id)

    calls: list[str] = []
    _patch_renderers(monkeypatch, calls)

    seen_names: list[str] = []

    class NamingDriveClient(FakeDriveClient):
        async def upload_new(self, *, name, parent_id, content, mime_type, app_properties=None):
            seen_names.append(name)
            return await super().upload_new(
                name=name, parent_id=parent_id, content=content, mime_type=mime_type, app_properties=app_properties
            )

    client = NamingDriveClient(calls)
    _patch_drive_client(monkeypatch, client)

    await report_delivery.deliver_report_to_drive(
        db, tenant_id=tenant.id, report_id=report.id, actor_type="user", actor_id=user.id, period_key="2026-09-08"
    )

    assert "Inventory Aging Weekly — 2026-09-08.pdf" in seen_names
    assert "Inventory Aging Weekly — 2026-09-08.xlsx" in seen_names


async def test_system_actor_writes_null_actor_id_audit(db, monkeypatch):
    """A6's headless compose script delivers as ``actor_type='system'`` — the audit
    row's actor_id is NULL, matching the house convention for cron/script actors."""
    tenant = await create_test_tenant(db, name="SystemActor")
    user, _ = await create_test_user(db, tenant)  # only for _seed_report's created_by
    await set_tenant_context(db, str(tenant.id))
    report = await _seed_report(db, tenant, user)
    await _add_sheets_connector(db, tenant.id)

    calls: list[str] = []
    _patch_renderers(monkeypatch, calls)
    client = FakeDriveClient(calls)
    _patch_drive_client(monkeypatch, client)

    await report_delivery.deliver_report_to_drive(
        db, tenant_id=tenant.id, report_id=report.id, actor_type="system", actor_id=None, period_key="2026-09-08"
    )

    await set_tenant_context(db, str(tenant.id))
    row = (
        await db.execute(
            text(
                "SELECT actor_id, actor_type FROM audit_events "
                "WHERE action='report.delivery.completed' AND resource_id=:rid"
            ),
            {"rid": str(report.id)},
        )
    ).first()
    assert row is not None
    assert row[0] is None
    assert row[1] == "system"
